from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "output"
BASE_NAME = "MFEC_PM_MotifPRDDBN_Q2_2025_20260109"
DOCX_PATH = OUTPUT_DIR / f"{BASE_NAME}_PM_Report.docx"


REQUIRED_TEXT = [
    "Table of Contents",
    "Defect Classification",
    "Suggestion Summary",
    "1. General Project Information",
    "2. Oracle minimum requirement",
    "2.1 Checking server machine specification",
    "2.2 Compare to Oracle requirement",
    "2.3 User's environment",
    "3. System Checklist",
    "4. Database Information",
    "4.1 Database Configuration.",
    "4.2 Database Parameter",
    "4.4 Database Files",
    "4.5 Temporary Files",
    "4.6 Redo Log File",
    "4.7 Control File",
    "4.8 Database Patch",
    "5. RDBMS Performance",
    "5.1 Performance Review",
    "5.2 Database Growth Rate",
    "5.3 Performance Analysis",
    "6. Tablespace Free Space",
    "7. Default tablespace and temporary tablespace",
    "8. Database Registry",
    "APPENDIX A",
    "APPENDIX B",
    "APPENDIX C",
    "APPENDIX D",
    "APPENDIX E",
    "APPENDIX F",
]


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def collect_docx_text(docx_path: Path) -> str:
    document = Document(docx_path)
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    chunks.append(text)

    return "\n".join(chunks)


def main() -> int:
    if not DOCX_PATH.exists():
        fail(f"DOCX not found: {DOCX_PATH}")

    text = collect_docx_text(DOCX_PATH)
    missing = [item for item in REQUIRED_TEXT if item not in text]

    if missing:
        for item in missing:
            print(f"MISSING: {item}")
        fail(f"DOCX structure validation failed with {len(missing)} missing section(s).")

    for item in REQUIRED_TEXT:
        print(f"OK: found {item}")

    print("DOCX structure validation passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

