from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "output"
BASE_NAME = "MFEC_PM_MotifPRDDBN_Q2_2025_20260109"
DOCX_PATH = OUTPUT_DIR / f"{BASE_NAME}_PM_Report.docx"
JSON_PATH = OUTPUT_DIR / f"{BASE_NAME}_parsed_data.json"

EXPECTED_CHARTS = [
    "database_growth_rate.png",
    "database_pga_advisory.png",
    "database_sga_advisory.png",
    "os_runq_perf.png",
    "os_cpu_util_all.png",
    "os_cpu_util_system.png",
    "os_cpu_util_user.png",
    "os_mem_free.png",
    "os_swap_used.png",
    "os_page_in.png",
    "os_page_out.png",
    "os_cpu_wait_io.png",
    "os_disk_service_time.png",
    "os_cpu_perf.png",
    "os_disk_perf.png",
    "os_mem_perf.png",
]


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def warn(message: str) -> None:
    print(f"WARN: {message}")


def ok(message: str) -> None:
    print(f"OK: {message}")


def collect_docx_text(document: Document) -> str:
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    chunks.append(text)

    return "\n".join(chunks)


def count_docx_media(docx_path: Path) -> int:
    with zipfile.ZipFile(docx_path) as archive:
        return len([name for name in archive.namelist() if name.startswith("word/media/")])


def assert_contains(label: str, text: str, expected: str) -> None:
    if expected not in text:
        fail(f"{label} not found in DOCX text: {expected}")
    ok(f"{label} found in DOCX text")


def main() -> int:
    if not DOCX_PATH.exists():
        fail(f"DOCX not found: {DOCX_PATH}")
    if not JSON_PATH.exists():
        fail(f"Parsed JSON not found: {JSON_PATH}")

    with JSON_PATH.open("r", encoding="utf-8") as file:
        data = json.load(file)

    document = Document(DOCX_PATH)
    text = collect_docx_text(document)

    table_count = len(document.tables)
    paragraph_count = len([paragraph for paragraph in document.paragraphs if paragraph.text.strip()])
    inline_shape_count = len(document.inline_shapes)
    media_count = count_docx_media(DOCX_PATH)

    if DOCX_PATH.stat().st_size < 300_000:
        fail(f"DOCX size is suspiciously small: {DOCX_PATH.stat().st_size} bytes")
    ok(f"DOCX size is {DOCX_PATH.stat().st_size} bytes")

    if paragraph_count < 120:
        fail(f"Too few non-empty paragraphs: {paragraph_count}")
    ok(f"Non-empty paragraphs: {paragraph_count}")

    if table_count < 35:
        fail(f"Too few tables for the expected PM report structure: {table_count}")
    ok(f"Tables: {table_count}")

    if inline_shape_count < 6:
        fail(f"Expected at least 6 embedded images/charts, found {inline_shape_count}")
    ok(f"Inline images/charts: {inline_shape_count}")

    if media_count < 6:
        fail(f"Expected at least 6 media files inside DOCX, found {media_count}")
    ok(f"DOCX media files: {media_count}")

    for chart_name in EXPECTED_CHARTS:
        chart_path = OUTPUT_DIR / chart_name
        if not chart_path.exists():
            fail(f"Expected chart file not found: {chart_path}")
        if chart_path.stat().st_size < 10_000:
            fail(f"Chart file is suspiciously small: {chart_path}")
        ok(f"Chart exists: {chart_name} ({chart_path.stat().st_size} bytes)")

    os_info = data.get("os_info", {})
    assert_contains("Hostname", text, os_info.get("hostname", ""))
    assert_contains("IP address", text, os_info.get("ip_address", ""))
    assert_contains("OS version", text, os_info.get("os_version", ""))
    assert_contains("Oracle home", text, os_info.get("oracle_home", ""))

    for registry_id in ("CATALOG", "CATPROC", "JAVAVM", "XML", "XDB", "SDO", "XOQ"):
        assert_contains(f"Registry component {registry_id}", text, registry_id)

    backup_info = data.get("backup_info", {})
    if backup_info.get("type") == "rman_log":
        fail("Appendix E selected RMAN fallback instead of Datapump backup evidence")
    assert_contains("Appendix E Datapump reference", text, "Datapump")
    assert_contains("Appendix F run queue section", text, "Process Run Queue")
    assert_contains("Appendix F CPU threshold table", text, "Very High (100%)")
    assert_contains("Appendix F physical memory section", text, "Physical Memory")
    assert_contains("Appendix F disk wait section", text, "CPU Wait I/O")

    placeholder_patterns = {
        "N/A": r"\bN/A\b",
        "None": r"\bNone\b",
        "TBD": r"\bTBD\b",
        "placeholder": r"placeholder",
        "not found": r"not found",
    }
    placeholder_hits = {
        label: len(re.findall(pattern, text, flags=re.IGNORECASE))
        for label, pattern in placeholder_patterns.items()
    }

    if placeholder_hits["TBD"] or placeholder_hits["placeholder"]:
        fail(f"Draft placeholders found: {placeholder_hits}")

    if placeholder_hits["N/A"] > 10:
        warn(f"High N/A count: {placeholder_hits['N/A']}")
    else:
        ok(f"N/A count acceptable: {placeholder_hits['N/A']}")

    if placeholder_hits["None"] > 15:
        warn(f"High None count: {placeholder_hits['None']}")
    else:
        ok(f"None count acceptable: {placeholder_hits['None']}")

    if placeholder_hits["not found"] > 5:
        warn(f"High 'not found' count: {placeholder_hits['not found']}")
    else:
        ok(f"'not found' count acceptable: {placeholder_hits['not found']}")

    print("Report quality/content validation passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
