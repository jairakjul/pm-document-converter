from pathlib import Path
from datetime import datetime
import json
import re
import sys

from docx import Document
from docx.shared import RGBColor, Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from matplotlib.patches import Circle, Arc
import numpy as np


BLUE = "0000FF"
HEADER_BLUE = "0070C0"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "D9D9D9"
YELLOW = "FFFF00"
RED = "FF0000"
WARNING_YELLOW = "FFFF00"
OK_GREEN = "92D050"
WHITE = "FFFFFF"
BLACK = "000000"
PROJECT_ROOT = Path(__file__).resolve().parents[1]
_TRANSLATION_CACHE: dict[str, dict[str, str]] = {}


def _translation_roots() -> list[Path]:
    roots = [PROJECT_ROOT / "translations"]
    bundled_root = Path(getattr(sys, "_MEIPASS", PROJECT_ROOT)) / "translations"
    if bundled_root not in roots:
        roots.insert(0, bundled_root)
    return roots


def _load_translations(language_code: str) -> dict[str, str]:
    code = "th" if language_code == "th" else "en"
    if code in _TRANSLATION_CACHE:
        return _TRANSLATION_CACHE[code]

    translations: dict[str, str] = {}
    for root in _translation_roots():
        path = root / f"{code}.json"
        if not path.exists():
            continue
        try:
            loaded = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            loaded = {}
        if isinstance(loaded, dict):
            translations = {str(key): str(value) for key, value in loaded.items()}
            break

    _TRANSLATION_CACHE[code] = translations
    return translations


def _report_language(data: dict) -> str:
    language = str(data.get("gui_metadata", {}).get("language", "English")).strip().lower()
    return "th" if language.startswith(("thai", "th")) else "en"


def _t(data: dict, text: str) -> str:
    return _load_translations(_report_language(data)).get(text, text)


def _report_type_label(data: dict, report_type: str) -> str:
    return _t(data, report_type)


def _footer_last_update(gui_meta: dict) -> str:
    candidates = [str(gui_meta.get("report_date", "")).strip()]
    change_records = gui_meta.get("change_records", [])
    if change_records:
        candidates.append(str(change_records[0].get("date", "")).strip())

    for candidate in candidates:
        if not candidate:
            continue
        for fmt in (
            "%d-%b-%Y",
            "%d %b %Y",
            "%d-%B-%Y",
            "%d %B %Y",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%m/%d/%Y",
            "%B %d, %Y",
        ):
            try:
                dt = datetime.strptime(candidate, fmt)
                return f"{dt.month}/{dt.day}/{dt.year}"
            except ValueError:
                continue
        return candidate

    today = datetime.now()
    return f"{today.month}/{today.day}/{today.year}"
MFEC_LOGO_PATHS = [
    PROJECT_ROOT / "assets" / "mfec_logo.png",
    Path(r"C:\Users\HP\Pictures\Picture2.png"),
    Path(r"C:\Users\HP\Pictures\Picture1.png"),
]
COVER_BANNER_PATHS = [
    PROJECT_ROOT / "assets" / "cover_banner.jpg",
    Path(r"C:\Users\HP\Pictures\Picture1.jpg"),
]


# =========================================================
# Basic Word helpers
# =========================================================

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def get_mfec_logo_path():
    for path in MFEC_LOGO_PATHS:
        if path.exists():
            return path
    return None


def get_cover_banner_path():
    for path in COVER_BANNER_PATHS:
        if path.exists():
            return path
    return None


def set_cell_text(cell, text, bold=False, color=BLACK, size=8, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]

    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(str(text))
    run.font.name = "Tahoma"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="000000", size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_monospaced_box(doc, text, font_size=7):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, color="000000", size="6")

    cell = table.cell(0, 0)
    cell.text = ""
    for index, line in enumerate(str(text).splitlines()):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(BLACK)

    return table


def set_doc_layout(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.different_first_page_header_footer = True

    styles = doc.styles
    styles["Normal"].font.name = "Tahoma"
    styles["Normal"].font.size = Pt(8)


def add_report_footer(doc, data):
    gui_meta = data.get("gui_metadata", {})
    report_date = _footer_last_update(gui_meta)

    confidential_text = "MFEC & SLI Confidential"

    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False

    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_top_border(paragraph)

    table = footer.add_table(rows=2, cols=3, width=Inches(7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_column_widths(table, [Inches(2.2), Inches(2.8), Inches(2.2)])

    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = left.add_run(f"Last Update.\n{report_date}")
    run.font.name = "Tahoma"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    middle = table.cell(0, 1).paragraphs[0]
    middle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = middle.add_run(confidential_text)
    run.font.name = "Tahoma"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    right = table.cell(0, 2).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right.add_run("Page ")
    run.font.name = "Tahoma"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    add_page_number_field(right)

    swatch_cell = table.cell(1, 1)
    swatch_paragraph = swatch_cell.paragraphs[0]
    swatch_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for color in ("A6A1D6", "ADA5F7", "B9CBFA", "D8ECFB"):
        swatch = swatch_paragraph.add_run("  ")
        swatch.font.size = Pt(7)
        swatch.font.highlight_color = None
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), color)
        swatch._r.get_or_add_rPr().append(shd)
        spacer = swatch_paragraph.add_run("  ")
        spacer.font.size = Pt(7)

    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)


def add_report_header(doc, data):
    gui_meta = data.get("gui_metadata", {})
    project_no = str(gui_meta.get("project_no", "")).strip()

    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    table = header.add_table(rows=1, cols=2, width=Inches(7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_column_widths(table, [Inches(6.1), Inches(1.1)])

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_margins(left_cell, top=0, bottom=0, left=0, right=0)
    set_cell_margins(right_cell, top=0, bottom=0, left=0, right=0)

    left_p = left_cell.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = left_p.add_run(project_no)
    left.font.name = "Tahoma"
    left.font.size = Pt(9)
    left.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    right_p = right_cell.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo_path = get_mfec_logo_path()
    if logo_path:
        right_p.add_run().add_picture(str(logo_path), width=Inches(0.82))
    else:
        logo = right_p.add_run("MFEC")
        logo.font.name = "Tahoma"
        logo.font.size = Pt(18)
        logo.font.bold = True
        logo.font.color.rgb = RGBColor(0x9A, 0x9A, 0x9A)

    line = header.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    line.paragraph_format.space_after = Pt(0)
    add_bottom_border(line, color="808080", size="4", space="1")


def add_bottom_border(paragraph, color=BLUE, size="6", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_top_border(paragraph, color="808080", size="4", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), size)
    top.set(qn("w:space"), space)
    top.set(qn("w:color"), color)
    p_bdr.append(top)
    p_pr.append(p_bdr)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)
    run.font.name = "Tahoma"
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return run


def set_paragraph_outline_level(paragraph, level):
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = r'TOC \o "1-3" \h \z \u'

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update field to refresh table of contents."

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(placeholder)
    run._r.append(fld_char_end)
    run.font.name = "Tahoma"
    run.font.size = Pt(8)
    return run


def set_update_fields_on_open(doc):
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def clear_cell(cell):
    cell._tc.clear_content()
    cell._tc.append(OxmlElement("w:p"))


def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for edge, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        margin = tc_mar.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_table_column_widths(table, widths):
    for row in table.rows:
        for index, width in enumerate(widths):
            if index < len(row.cells):
                row.cells[index].width = width


def set_fixed_table_grid(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = table._tbl.tblGrid
    if old_grid is not None:
        table._tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width.inches * 1440)))
        grid.append(col)
    table._tbl.insert(0, grid)
    set_table_column_widths(table, widths)


def add_nested_table(cell, headers, rows, header_fill=LIGHT_GRAY, header_color=BLACK, font_size=7, column_widths=None):
    clear_cell(cell)
    table = cell.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    if column_widths:
        set_fixed_table_grid(table, column_widths)

    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], header_fill)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=header_color, size=font_size)
        table.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(table.rows[0].cells[i], top=25, bottom=25, left=40, right=40)

    for row in rows:
        cells = table.add_row().cells
        for i, header in enumerate(headers):
            align = "left" if i == 0 else "right"
            set_cell_shading(cells[i], WHITE)
            set_cell_text(cells[i], row.get(header, ""), size=font_size, align=align)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[i], top=20, bottom=20, left=35, right=35)

    if column_widths:
        set_fixed_table_grid(table, column_widths)

    return table


def add_heading1(doc, text):
    p = doc.add_paragraph()
    set_paragraph_outline_level(p, 0)
    run = p.add_run(text)
    run.font.name = "Tahoma"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p_format = p.paragraph_format
    p_format.space_before = Pt(6)
    p_format.space_after = Pt(2)

    # blue underline
    border_p = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    border_p.append(p_bdr)

    return p


def add_heading2(doc, text):
    p = doc.add_paragraph()
    set_paragraph_outline_level(p, 1)
    run = p.add_run(text)
    run.font.name = "Tahoma"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)

    return p


def add_heading3(doc, text):
    p = doc.add_paragraph()
    set_paragraph_outline_level(p, 2)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLACK)

    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)

    return p


def add_normal(doc, text, size=8, bold=False, highlight_status_words=True):
    p = doc.add_paragraph()
    text = str(text)
    if highlight_status_words and re.search(r"\b(in)?sufficient\b", text, flags=re.IGNORECASE):
        parts = re.split(r"\b(insufficient|sufficient)\b", text, flags=re.IGNORECASE)
        for part in parts:
            if not part:
                continue
            run = p.add_run(part)
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            lower_part = part.lower()
            if lower_part == "sufficient":
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            elif lower_part == "insufficient":
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)
            else:
                run.font.color.rgb = RGBColor.from_string(BLACK)
    else:
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(BLACK)
    return p


def add_small_table(doc, headers, rows, header_fill=LIGHT_GRAY, highlight_rule=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_borders(table)

    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], header_fill)
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8)

    for row in rows:
        cells = table.add_row().cells
        for i, h in enumerate(headers):
            value = row.get(h, "")
            fill = WHITE

            if highlight_rule and highlight_rule(row, h, value):
                fill = YELLOW

            set_cell_shading(cells[i], fill)
            set_cell_text(cells[i], value, size=8)

    return table


def _safe_float(value, default=0.0):
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return default


def _build_tablespace_capacity_rows(data):
    tablespaces = {}

    def ensure_entry(name):
        entry = tablespaces.get(name)
        if entry is None:
            entry = {
                "name": name,
                "allocated_mb": 0.0,
                "max_mb": 0.0,
                "used_mb": 0.0,
                "free_of_max_mb": 0.0,
                "free_pct_of_max": 0.0,
                "status": "OK",
                "is_temp": False,
            }
            tablespaces[name] = entry
        return entry

    for source_key, is_temp in (("datafiles", False), ("tempfiles", True)):
        for file_row in data.get(source_key, []):
            name = str(file_row.get("tbs_name", "")).strip()
            if not name:
                continue

            allocated_mb = _safe_float(file_row.get("size_mb", 0))
            configured_max_mb = _safe_float(file_row.get("max_mb", 0))
            effective_max_mb = configured_max_mb if configured_max_mb > 0 else allocated_mb

            entry = ensure_entry(name)
            entry["allocated_mb"] += allocated_mb
            entry["max_mb"] += effective_max_mb
            entry["is_temp"] = entry["is_temp"] or is_temp

    tablespace_free = data.get("tablespace_free", {})
    for name, entry in tablespaces.items():
        if entry["is_temp"]:
            used_mb = 0.0
        else:
            free_pct_of_allocated = _safe_float(tablespace_free.get(name, {}).get("pct_free", 0))
            used_mb = entry["allocated_mb"] * max(0.0, 1 - (free_pct_of_allocated / 100.0))
            if not tablespace_free.get(name):
                used_mb = entry["allocated_mb"]

        free_of_max_mb = max(entry["max_mb"] - used_mb, 0.0)
        free_pct_of_max = (free_of_max_mb / entry["max_mb"] * 100.0) if entry["max_mb"] else 0.0

        entry["used_mb"] = used_mb
        entry["free_of_max_mb"] = free_of_max_mb
        entry["free_pct_of_max"] = free_pct_of_max
        if free_pct_of_max < 10:
            entry["status"] = "Critical"
        elif free_pct_of_max < 20:
            entry["status"] = "Warning"

    preferred_order = [
        "ICOM_TBS",
        "TSLI_COMP_VERF_TBS",
        "STATSPACK",
        "SYSAUX",
        "SYSTEM",
        "USERS",
        "AUDIT_TBS",
        "UNDOTBS1",
        "MONITOR_TBS",
        "TSLI_COMP_VERF_TMP",
        "ICOM_TMP",
        "ICOM_HOLD_TBS",
        "TMP",
    ]
    preferred_rank = {name: index for index, name in enumerate(preferred_order)}

    return sorted(
        tablespaces.values(),
        key=lambda row: (
            round(row["free_pct_of_max"], 2),
            preferred_rank.get(row["name"], 999),
            row["name"],
        ),
    )


def _cover_report_date(gui_meta, source_zip):
    change_records = gui_meta.get("change_records", [])
    candidate = str(gui_meta.get("report_date", "")).strip()

    if candidate:
        for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y", "%B %d, %Y"):
            try:
                dt = datetime.strptime(candidate, fmt)
                return dt.strftime("%B %d, %Y").replace(" 0", " ")
            except ValueError:
                continue

    if change_records:
        candidate = str(change_records[0].get("date", "")).strip()

    if not candidate:
        match = re.search(r"(20\d{6})", source_zip.stem)
        if match:
            try:
                candidate = datetime.strptime(match.group(1), "%Y%m%d").strftime("%d-%b-%Y")
            except ValueError:
                candidate = ""

    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            dt = datetime.strptime(candidate, fmt)
            return dt.strftime("%B %d, %Y").replace(" 0", " ")
        except ValueError:
            continue

    return datetime.now().strftime("%B %d, %Y").replace(" 0", " ")


def create_cover_banner(output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "cover_banner.png"

    width, height = 1400, 360
    x = np.linspace(0.0, 1.0, width)
    y = np.linspace(0.0, 1.0, height)
    xx, yy = np.meshgrid(x, y)

    base = np.zeros((height, width, 3))
    base[:, :, 0] = 0.10 + 0.12 * (1 - xx) + 0.04 * yy
    base[:, :, 1] = 0.30 + 0.35 * (1 - np.abs(yy - 0.5)) + 0.10 * xx
    base[:, :, 2] = 0.65 + 0.25 * (1 - xx * 0.6)

    fig, ax = plt.subplots(figsize=(9.2, 2.5), dpi=150)
    ax.imshow(base, extent=[0, 1, 0, 1], aspect="auto")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.axhspan(0.88, 1.0, color="#1100EE")

    globe_fill = Circle((0.16, 0.48), 0.28, facecolor=(0.86, 0.96, 1.0, 0.70), edgecolor="#0B2BFF", linewidth=2)
    globe_edge = Circle((0.16, 0.48), 0.28, facecolor=(0, 0, 0, 0), edgecolor="#0038FF", linewidth=2.2)
    ax.add_patch(globe_fill)
    ax.add_patch(globe_edge)

    for offset in (-0.17, 0.0, 0.17):
        ax.add_patch(Arc((0.16, 0.48), 0.56, 0.56 * max(0.18, 1 - abs(offset) * 1.8), angle=0, theta1=0, theta2=360, edgecolor="#1A4CFF", linewidth=1.0, alpha=0.75))
    for offset in (-55, -20, 15, 50):
        ax.add_patch(Arc((0.16, 0.48), 0.56, 0.56, angle=offset, theta1=35, theta2=145, edgecolor="#1A4CFF", linewidth=0.9, alpha=0.65))

    continent_x = [0.05, 0.08, 0.10, 0.13, 0.15, 0.17, 0.19, 0.21, 0.20, 0.17, 0.14, 0.10]
    continent_y = [0.65, 0.72, 0.70, 0.76, 0.70, 0.66, 0.60, 0.54, 0.46, 0.40, 0.44, 0.55]
    ax.fill(continent_x, continent_y, color=(0.95, 1.0, 1.0, 0.9))
    ax.fill([0.09, 0.11, 0.13, 0.12, 0.10], [0.39, 0.34, 0.28, 0.23, 0.18], color=(0.95, 1.0, 1.0, 0.85))

    for x0 in np.linspace(0.36, 0.92, 7):
        ax.plot([x0 - 0.08, x0 + 0.08], [0.30, 0.76], color=(0.70, 0.94, 1.0, 0.30), linewidth=4)

    points = [(0.54, 0.60), (0.62, 0.66), (0.72, 0.54), (0.82, 0.70), (0.90, 0.52), (0.68, 0.38)]
    for x0, y0 in points:
        ax.scatter([x0], [y0], s=48, color=(1, 1, 1, 0.85), edgecolors=(0.72, 0.95, 1.0, 0.55), linewidths=1)
    for (x1, y1), (x2, y2) in zip(points[:-1], points[1:]):
        ax.plot([x1, x2], [y1, y2], color=(0.82, 0.97, 1.0, 0.55), linewidth=2)

    digits = ["0101", "1010", "0011", "1100", "0110", "1001"]
    coords = [(0.63, 0.82), (0.73, 0.76), (0.83, 0.80), (0.70, 0.22), (0.88, 0.30), (0.56, 0.18)]
    for (x0, y0), text in zip(coords, digits):
        ax.text(x0, y0, text, color=(0.85, 0.98, 1.0, 0.42), fontsize=28, fontweight="bold", rotation=18)

    plt.tight_layout(pad=0)
    plt.savefig(output_path, bbox_inches="tight", pad_inches=0)
    plt.close(fig)
    return output_path


def add_cover_page(doc, data, source_zip, output_dir):
    gui_meta = data.get("gui_metadata", {})
    db_name = gui_meta.get("db_name") or data.get("instance_name") or data.get("db_parameters", {}).get("db_name", "")
    report_type = gui_meta.get("report_type") or "Preventive Maintenance"
    project_name = str(gui_meta.get("project_name", "")).strip()
    customer_name = (
        str(gui_meta.get("customer_full", "")).strip()
        or str(gui_meta.get("customer_name", "")).strip()
        or next((row.get("name", "").strip() for row in gui_meta.get("customers", []) if row.get("name", "").strip()), "")
        or str(gui_meta.get("customer_abbrev", "")).strip()
    )
    report_date = _cover_report_date(gui_meta, source_zip)

    title_text = f"{_report_type_label(data, report_type)} {db_name}".strip()
    if not title_text.strip():
        title_text = project_name or _report_type_label(data, report_type)

    confidential = doc.add_paragraph()
    confidential.alignment = WD_ALIGN_PARAGRAPH.CENTER
    confidential.paragraph_format.space_before = Pt(22)
    confidential.paragraph_format.space_after = Pt(6)
    confidential_run = confidential.add_run("MFEC & SLI Confidential")
    confidential_run.font.name = "Arial"
    confidential_run.font.size = Pt(14)
    confidential_run.font.color.rgb = RGBColor.from_string(BLUE)

    blue_bar = doc.add_table(rows=1, cols=1)
    blue_bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    blue_bar.autofit = False
    set_table_column_widths(blue_bar, [Inches(7.3)])
    blue_cell = blue_bar.cell(0, 0)
    set_cell_margins(blue_cell, top=35, bottom=35, left=0, right=0)
    set_cell_shading(blue_cell, BLUE)
    blue_cell.text = ""

    banner_path = get_cover_banner_path()
    if banner_path:
        banner = doc.add_paragraph()
        banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        banner.paragraph_format.space_before = Pt(0)
        banner.paragraph_format.space_after = Pt(20)
        banner.add_run().add_picture(str(banner_path), width=Inches(7.3), height=Inches(1.85))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title_run = title.add_run(title_text)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(25)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string("FF6600")

    if customer_name:
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        subtitle.paragraph_format.space_after = Pt(14)
        subtitle_run = subtitle.add_run(f"{_t(data, 'For')} {customer_name}")
        subtitle_run.font.name = "Arial"
        subtitle_run.font.size = Pt(22)
        subtitle_run.font.color.rgb = RGBColor.from_string("FF0000")

    divider = doc.add_paragraph()
    divider.paragraph_format.space_before = Pt(8)
    divider.paragraph_format.space_after = Pt(10)
    add_bottom_border(divider, color=BLUE, size="6", space="1")

    date_paragraph = doc.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_run = date_paragraph.add_run(f"{_t(data, 'Date')}: {report_date}")
    date_run.font.name = "Arial"
    date_run.font.size = Pt(17)
    date_run.font.italic = True
    date_run.font.color.rgb = RGBColor.from_string(BLUE)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(42)
    spacer.paragraph_format.space_after = Pt(0)
    dpm_table = doc.add_table(rows=1, cols=2)
    dpm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dpm_table.autofit = False
    set_table_column_widths(dpm_table, [Inches(3.1), Inches(3.8)])

    swatch_cell = dpm_table.cell(0, 0)
    set_cell_margins(swatch_cell, top=0, bottom=0, left=24, right=0)
    swatch_table = swatch_cell.add_table(rows=3, cols=5)
    swatch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    swatch_table.autofit = False
    set_table_column_widths(
        swatch_table,
        [Inches(0.55), Inches(0.18), Inches(0.55), Inches(0.18), Inches(0.55)],
    )
    set_table_borders(swatch_table, color="FFFFFF", size="0")
    swatch_rows = [
        ["69B7B4", None, "7F98F2", None, "D9D9D9"],
        [None, None, None, None, None],
        ["6A67B6", None, "BBD9F4", None, "BFBFBF"],
    ]
    for row_index, row_colors in enumerate(swatch_rows):
        row = swatch_table.rows[row_index]
        row.height = Inches(0.55 if row_index != 1 else 0.18)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for cell, color in zip(row.cells, row_colors):
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_text(cell, "", size=1, align="center")
            set_cell_shading(cell, color or WHITE)

    dpm_cell = dpm_table.cell(0, 1)
    set_cell_margins(dpm_cell, top=52, bottom=0, left=0, right=0)
    dpm_p = dpm_cell.paragraphs[0]
    dpm_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    dpm_run = dpm_p.add_run("By: Digital Platform and Management\n(DPM)\nMFEC Public Company Limited")
    dpm_run.font.name = "Arial"
    dpm_run.font.size = Pt(10)
    dpm_run.font.bold = True

    bottom_spacer = doc.add_paragraph()
    bottom_spacer.paragraph_format.space_before = Pt(70)
    bottom_spacer.paragraph_format.space_after = Pt(0)
    cover_table = doc.add_table(rows=1, cols=3)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    set_table_column_widths(cover_table, [Inches(1.7), Inches(2.65), Inches(2.65)])

    logo_cell = cover_table.cell(0, 0)
    set_cell_margins(logo_cell, top=0, bottom=0, left=0, right=0)
    logo_paragraph = logo_cell.paragraphs[0]
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_path = get_mfec_logo_path()
    if logo_path:
        logo_paragraph.add_run().add_picture(str(logo_path), width=Inches(1.55))
    else:
        run = logo_paragraph.add_run("MFEC")
        run.font.name = "Arial"
        run.font.size = Pt(32)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x40, 0x45, 0x4A)

    office_cell = cover_table.cell(0, 1)
    set_cell_margins(office_cell, top=0, bottom=0, left=15, right=15)
    office_paragraph = office_cell.paragraphs[0]
    office_run = office_paragraph.add_run(
        "Head Office:\n"
        "349 SJ Infinite One Business Complex,11th\n"
        "Floor,Vibhavadi Rangsit Rd, Chompol, Chatuchak,\n"
        "Bangkok 10900. Thailand Tel: +66 (0) 2821-7999\n"
        "www.mfec.co.th"
    )
    office_run.font.name = "Arial"
    office_run.font.size = Pt(7)
    office_run.font.bold = False

    dev_cell = cover_table.cell(0, 2)
    set_cell_margins(dev_cell, top=0, bottom=0, left=15, right=0)
    dev_paragraph = dev_cell.paragraphs[0]
    dev_run = dev_paragraph.add_run(
        "Development Center:\n"
        "199 S-Oasis 21 Floor, Vibhavadi-Rangsit Rd.,\n"
        "Chompol,Chatuchak, Bangkok 10900. Thailand"
    )
    dev_run.font.name = "Arial"
    dev_run.font.size = Pt(7)

    doc.add_page_break()


# =========================================================
# 0. Manual TOC
# =========================================================

def add_table_of_contents(doc, data):
    """Add a real Word TOC field that updates from paragraph outline levels."""

    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.autofit = False
    set_table_borders(title_table, size="6")
    set_table_column_widths(title_table, [Inches(4.2)])
    title_cell = title_table.cell(0, 0)
    set_cell_margins(title_cell, top=10, bottom=10, left=20, right=20)
    set_cell_text(title_cell, _t(data, "Table of Contents"), bold=True, color=BLUE, size=14, align="center")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)

    toc_paragraph = doc.add_paragraph()
    toc_paragraph.paragraph_format.space_before = Pt(0)
    toc_paragraph.paragraph_format.space_after = Pt(0)
    add_toc_field(toc_paragraph)

    doc.add_page_break()
    return

    # Define TOC entries:  (title, page_number, level)
    # level 0 = main section,  level 1 = sub-section,  level 2 = sub-sub-section
    toc_rows = [
        ("Defect Classification", 4, 0),
        ("Suggestion Summary", 5, 0),
        ("1. General Project Information", 6, 0),
        ("1.1 Project Information", 6, 1),
        ("1.2 Customer Contact Information", 6, 1),
        ("1.3 MFEC Engineers' Information", 6, 1),
        ("1.4 Change Record", 6, 1),
        ("1.5 Reviewers", 6, 1),
        ("2. Oracle minimum requirement", 7, 0),
        ("2.1 Checking server machine specification", 7, 1),
        ("2.2 Compare to Oracle requirement for MotifPRDDBN", 7, 1),
        ("2.3 User's environment for MotifPRDDBN", 8, 1),
        ("3. System Checklist", 10, 0),
        ("3.1 Hardware configuration for MotifPRDDBN", 10, 1),
        ("3.2 Network configuration for MotifPRDDBN", 10, 1),
        ("3.3 Crontab information for MotifPRDDBN", 10, 1),
        ("4. Database Information", 11, 0),
        ("4.1 Database Configuration.", 11, 1),
        ("4.2 Database Parameter", 11, 1),
        ("4.3 Major Security Initialization Parameters", 13, 1),
        ("4.4 Database Files", 13, 1),
        ("4.5 Temporary Files", 14, 1),
        ("4.6 Redo Log File", 14, 1),
        ("4.7 Control File", 14, 1),
        ("4.8 Database Patch", 15, 1),
        ("5. RDBMS Performance", 16, 0),
        ("5.1 Performance Review", 16, 1),
        ("5.2 Database Growth Rate", 18, 1),
        ("5.3 Performance Analysis", 19, 1),
        ("6. Tablespace Free Space", 23, 0),
        ("6.1 Tablespace Free Space", 23, 1),
        ("7. Default tablespace and temporary tablespace", 24, 0),
        ("7.1 Default tablespace and temporary tablespace", 24, 1),
        ("8. Database Registry", 26, 0),
        ("8.1 Check Database Registry", 26, 1),
        ("APPENDIX A – Invalid Object", 27, 0),
        ("APPENDIX B – Information from alert log", 36, 0),
        ("APPENDIX C – SQL Statement should be to investigate", 37, 0),
        ("APPENDIX D – Operating system log", 38, 0),
        ("APPENDIX E – Backup Configuration", 39, 0),
        ("APPENDIX F – OS Performance Summary", 40, 0),
    ]

    # Style settings per level
    level_config = {
        0: {"font_size": Pt(6.7), "bold": False, "color": RGBColor(0x00, 0x00, 0x00), "indent_cm": 0},
        1: {"font_size": Pt(6.4), "bold": False, "color": RGBColor(0x00, 0x00, 0x00), "indent_cm": 0.35},
        2: {"font_size": Pt(6.2), "bold": False, "color": RGBColor(0x00, 0x00, 0x00), "indent_cm": 0.7},
    }

    right_tab_pos = 8600

    for title, page, level in toc_rows:
        cfg = level_config.get(level, level_config[0])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        # Set left indent
        p.paragraph_format.left_indent = Pt(cfg["indent_cm"] * 28.35)

        # Add a right-aligned tab stop with dot leader via XML
        pPr = p._p.get_or_add_pPr()
        tabs_elem = OxmlElement('w:tabs')
        tab_elem = OxmlElement('w:tab')
        tab_elem.set(qn('w:val'), 'right')
        tab_elem.set(qn('w:leader'), 'dot')
        tab_elem.set(qn('w:pos'), str(right_tab_pos))
        tabs_elem.append(tab_elem)
        pPr.append(tabs_elem)

        # Add the title text run
        r_title = p.add_run(title)
        r_title.font.name = "Arial"
        r_title.font.size = cfg["font_size"]
        r_title.font.bold = cfg["bold"]
        r_title.font.color.rgb = cfg["color"]

        # Add a tab character (this triggers the dot leader)
        r_tab = p.add_run("\t")
        r_tab.font.name = "Arial"
        r_tab.font.size = cfg["font_size"]

        # Add the page number
        r_page = p.add_run(str(page))
        r_page.font.name = "Arial"
        r_page.font.size = cfg["font_size"]
        r_page.font.bold = cfg["bold"]
        r_page.font.color.rgb = cfg["color"]

    doc.add_page_break()


# =========================================================
# Defect / Summary
# =========================================================

def add_defect_classification(doc, data):
    add_heading1(doc, _t(data, "Defect Classification"))
    add_normal(
        doc,
        "Each test will be either deemed as a pass or a severity level will be given using the following severity level classification.",
        size=8,
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_borders(table)

    headers = ["Severity\nLevel", "Definition", "Actions"]
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[i], h, bold=True, align="center")

    rows = [
        ("1", "Critical", "Chance of Database down, need immediate action", RED),
        ("2", "Warning", "Change of hindering some database function, doesn’t require immediate action", WARNING_YELLOW),
        ("3", "OK (Stable)", "Database can work normally. No action required.", OK_GREEN),
    ]

    for sev, definition, action, fill in rows:
        cells = table.add_row().cells
        for c in cells:
            set_cell_shading(c, fill)

        set_cell_text(cells[0], sev, align="center", bold=True)
        set_cell_text(cells[1], definition, align="center")
        set_cell_text(cells[2], action, align="left")

    doc.add_page_break()


def build_suggestion_summary(data):
    invalid_count = len(data.get("invalid_objects_list", []))
    patch_info = data.get("patch_info", [])
    tablespace_rows = _build_tablespace_capacity_rows(data)
    sec53 = data.get("section53", {})
    os_perf = data.get("os_perf", {})
    db_params = data.get("db_parameters", {})

    # --- Patch status ---
    patch_status = "OK"
    if patch_info:
        for p in patch_info:
            if "recommend" in str(p.get("suggestion", "")).lower():
                patch_status = "Critical"

    # --- Tablespace status ---
    tbs_status = "OK"
    for row in tablespace_rows:
        if row["status"] == "Critical":
            tbs_status = "Critical"
            break
        elif row["status"] == "Warning":
            tbs_status = "Warning"

    # --- Invalid objects ---
    invalid_status = "Warning" if invalid_count > 0 else "OK"

    # --- 4.2/4.3 audit_trail check ---
    audit_trail = str(db_params.get("audit_trail", "")).upper()
    param_status = "Warning" if "EXTENDED" in audit_trail or "DB" in audit_trail else "OK"

    # --- 4.4 Datafile inc_mb check ---
    df_status = "OK"
    for f in data.get("datafiles", []):
        try:
            inc = float(f.get("inc_mb", "0"))
            if 0 < inc < 5:
                df_status = "Warning"
                break
        except (ValueError, TypeError):
            pass

    # --- 5.1 Archive mode ---
    archive_mode = data.get("archive_mode", "")
    review_status = "OK" if archive_mode == "ARCHIVELOG" else "Warning"

    # --- 5.3 CPU Section (Parse CPU to Parse Elapsd) ---
    cpu_status = "OK"
    inst_eff = sec53.get("instance_efficiency", {})
    try:
        parse_val = float(inst_eff.get("Parse CPU to Parse Elapsd %", "100"))
        if parse_val < 50.0:
            cpu_status = "Warning"
    except ValueError:
        pass

    # --- 5.3 SQL Section (open cursors / dblink) ---
    sql_status = "OK"
    cursors = sec53.get("open_cursors", {})
    try:
        pct_cursor = float(str(cursors.get("pct", "0")).replace("%", ""))
        if pct_cursor >= 95:
            sql_status = "Warning"
    except ValueError:
        pass
    top5 = sec53.get("top5_events", [])
    if any("dblink" in ev.lower() for ev in top5):
        sql_status = "Warning"

    # --- 5.3 Memory Section ---
    mem_status = "OK"

    # --- APPENDIX F OS Performance ---
    os_status = "OK"
    if os_perf.get("found"):
        cpu_trend = os_perf.get("cpu_trend", [])
        if cpu_trend:
            max_cpu = max([r.get("usr", 0) + r.get("sys", 0) for r in cpu_trend], default=0)
            if max_cpu >= 95:
                os_status = "Warning"
        runq = os_perf.get("runq_trend", [])
        if runq:
            max_runq = max((r.get("max", 0) for r in runq), default=0)
            try:
                cpu_count = int(data.get("os_info", {}).get("cpu_cores", "2"))
            except ValueError:
                cpu_count = 2
            if max_runq > cpu_count * 2:
                os_status = "Warning"

    return [
        ("2. Oracle Minimum Requirement", "OK", ""),
        ("3. System Checklist", "OK", ""),
        ("4. Database Information", "", ""),
        ("   4.1 Database Configuration", "OK", ""),
        ("   4.2 Database Parameter", param_status, "See more detail in 4.3" if param_status != "OK" else ""),
        ("   4.3 Major Security Initialization Parameter", param_status, "See more detail in 4.3" if param_status != "OK" else ""),
        ("   4.4 Database Files", df_status, "See more detail in 4.4" if df_status != "OK" else ""),
        ("   4.5 Temporary Files", "OK", ""),
        ("   4.6 Redo Log File", "OK", ""),
        ("   4.7 Control File", "OK", ""),
        ("   4.8 Database Patch", patch_status, "See more detail in 4.8" if patch_status != "OK" else ""),
        ("5. RDBMS Performance", "", ""),
        ("   5.1 Performance Review", review_status, "See more detail in 5.1" if review_status != "OK" else ""),
        ("   5.2 Database Growth Rate", "OK", ""),
        ("   5.3 Performance Analysis (CPU Section)", cpu_status, "See more detail in 5.3 (CPU Section)" if cpu_status != "OK" else ""),
        ("   5.3 Performance Analysis (SQL Section)", sql_status, "See more detail in 5.3 (SQL Section)" if sql_status != "OK" else ""),
        ("   5.3 Performance Analysis (Memory Section)", mem_status, ""),
        ("6. Tablespace Free Space", tbs_status, "See more detail in 6.1" if tbs_status != "OK" else ""),
        ("7. Default tablespace and temporary tablespace", "OK", ""),
        ("8. Database Registry", "OK", ""),
        ("APPENDIX A. Invalid Object", invalid_status, "See more detail in APPENDIX A" if invalid_status != "OK" else ""),
        ("APPENDIX B. Information from alert log", "OK", ""),
        ("APPENDIX C. SQL Statement should be to investigate", "OK", ""),
        ("APPENDIX D. Operating System Log", "OK", ""),
        ("APPENDIX E. Backup Configuration", "OK", ""),
        ("APPENDIX F. OS Performance Summary", os_status, "See more detail in APPENDIX F" if os_status != "OK" else ""),
    ]



def add_suggestion_summary(doc, data):
    add_heading1(doc, _t(data, "Suggestion Summary"))

    rows = build_suggestion_summary(data)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    set_table_borders(table)

    headers = ["Topic", "Critical", "Warning", "OK", "Remark"]
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], HEADER_BLUE)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, align="center")

    for topic, status, remark in rows:
        cells = table.add_row().cells

        for c in cells:
            set_cell_shading(c, WHITE)

        set_cell_text(cells[0], topic, size=7)
        set_cell_text(cells[1], "x" if status == "Critical" else "", align="center", size=7)
        set_cell_text(cells[2], "x" if status == "Warning" else "", align="center", size=7)
        set_cell_text(cells[3], "x" if status == "OK" else "", align="center", size=7)
        set_cell_text(cells[4], remark, size=7)

    doc.add_page_break()


# =========================================================
# Project Information
# =========================================================

def add_general_project_information(doc, data, source_zip):
    add_heading1(doc, _t(data, "1. General Project Information"))

    add_heading2(doc, _t(data, "1.1 Project Information"))

    gui_meta = data.get("gui_metadata", {})
    db_name = gui_meta.get("db_name") or data.get("instance_name") or data.get("db_parameters", {}).get("db_name", "")
    sales_list = gui_meta.get("sales", [])
    first_sale = sales_list[0] if sales_list else {}
    
    project_rows = [
        {"Field": "Project name", "Value": gui_meta.get("project_name", "MA Oracle 5 Years")},
        {"Field": "MFEC project no.", "Value": gui_meta.get("project_no", "")},
        {"Field": "MFEC sales name", "Value": gui_meta.get("sale", "")},
        {"Field": "Phone no.", "Value": first_sale.get("phone", "")},
        {"Field": "Email address", "Value": first_sale.get("email", "")},
        {"Field": "Database name", "Value": db_name},
        {"Field": "Source ZIP", "Value": source_zip.name},
    ]
    add_small_table(doc, ["Field", "Value"], project_rows)

    add_heading2(doc, _t(data, "1.2 Customer Contact Information"))
    cust_list = gui_meta.get("customers", [])
    if not cust_list:
        cust_list = [{"name": "", "phone": "", "email": ""}]
        
    cust_rows = [{"Name": c.get("name",""), "Phone no.": c.get("phone",""), "Email address": c.get("email","")} for c in cust_list]
    add_small_table(doc, ["Name", "Phone no.", "Email address"], cust_rows)

    add_heading2(doc, _t(data, "1.3 MFEC Engineers' Information"))
    eng_list = gui_meta.get("engineers", [])
    if not eng_list:
        eng_list = [{"name": "", "phone": "", "email": ""}]
        
    eng_rows = [{"Name": e.get("name",""), "Phone no.": e.get("phone",""), "Email address": e.get("email","")} for e in eng_list]
    add_small_table(doc, ["Name", "Phone no.", "Email address"], eng_rows)

    add_heading2(doc, _t(data, "1.4 Change Record"))
    cr_list = gui_meta.get("change_records", [])
    if not cr_list:
        cr_list = [{"date": datetime.now().strftime("%d-%b-%Y"), "author": "", "version": "1.0", "ref": "Initial Document"}]
        
    cr_rows = [{"Date": r.get("date",""), "Author": r.get("author",""), "Version": r.get("version",""), "Change Reference": r.get("ref","")} for r in cr_list]
    add_small_table(doc, ["Date", "Author", "Version", "Change Reference"], cr_rows)

    add_heading2(doc, _t(data, "1.5 Reviewers"))
    rev_list = gui_meta.get("reviewers", [])
    if not rev_list:
        rev_list = [{"date": "", "name": "", "position": ""}]
        
    rev_rows = [{"Date": r.get("date",""), "Name": r.get("name",""), "Position": r.get("position","")} for r in rev_list]
    add_small_table(doc, ["Date", "Name", "Position"], rev_rows)

    doc.add_page_break()


# =========================================================
# Section 2: Oracle Minimum Requirement
# =========================================================

def add_oracle_minimum_requirement(doc, data):
    os_info = data.get("os_info", {})

    add_heading1(doc, _t(data, "2. Oracle minimum requirement"))

    # ----- 2.1 Checking server machine specification -----
    add_heading2(doc, _t(data, "2.1 Checking server machine specification"))
    add_normal(doc, _t(data, "Server machine specification collected from the operating system:"))

    spec_rows = [
        {"Item": "Hostname", "Detail": os_info.get("hostname", "")},
        {"Item": "IP Address", "Detail": os_info.get("ip_address", "")},
        {"Item": "Subnet Mask", "Detail": os_info.get("subnet_mask", "")},
        {"Item": "OS Version", "Detail": os_info.get("os_version", "")},
        {"Item": "Kernel Version", "Detail": os_info.get("kernel_version", "")},
        {"Item": "Architecture", "Detail": os_info.get("architecture", "")},
        {"Item": "CPU Model", "Detail": os_info.get("cpu_model", "")},
        {"Item": "Number of CPU(s)", "Detail": os_info.get("cpu_count", "")},
        {"Item": "Core(s) per socket", "Detail": os_info.get("cores_per_socket", "")},
        {"Item": "Socket(s)", "Detail": os_info.get("sockets", "")},
        {"Item": "CPU MHz", "Detail": os_info.get("cpu_mhz", "")},
        {"Item": "CPU Cache Size", "Detail": os_info.get("cache_size", "")},
        {"Item": "Total RAM (GB)", "Detail": os_info.get("total_ram_gb", "")},
        {"Item": "Total Swap (MB)", "Detail": os_info.get("swap_total_mb", "")},
        {"Item": "Swap Used (MB)", "Detail": os_info.get("swap_used_mb", "")},
        {"Item": "/tmp Total Size", "Detail": os_info.get("tmp_size", "")},
        {"Item": "/tmp Available", "Detail": os_info.get("tmp_avail", "")},
        {"Item": "JDK/JRE Version", "Detail": os_info.get("jdk_version", "")},
    ]
    add_small_table(doc, ["Item", "Detail"], spec_rows)

    bundle_nodes = data.get("bundle_nodes", [])
    visible_nodes = [
        node for node in bundle_nodes
        if node.get("os_info", {}).get("hostname") or node.get("has_database_results") or node.get("has_registry")
    ]
    if len(visible_nodes) > 1:
        add_normal(doc, "Fast Assessment bundle sources detected:")
        node_rows = []
        for node in visible_nodes:
            node_os = node.get("os_info", {})
            node_rows.append(
                {
                    "Source ZIP": node.get("source", ""),
                    "Hostname": node_os.get("hostname", ""),
                    "IP Address": node_os.get("ip_address", ""),
                    "Oracle SID": node_os.get("oracle_sid", ""),
                    "Oracle Home": node_os.get("oracle_home", ""),
                    "DB Results": "Yes" if node.get("has_database_results") else "No",
                }
            )
        add_small_table(
            doc,
            ["Source ZIP", "Hostname", "IP Address", "Oracle SID", "Oracle Home", "DB Results"],
            node_rows,
        )

    # ----- 2.2 Compare to Oracle requirement -----
    hostname = os_info.get("hostname", "")
    compare_heading = f"{_t(data, '2.2 Compare to Oracle requirement')} {_t(data, 'For').lower()} {hostname}" if hostname else _t(data, "2.2 Compare to Oracle requirement")
    add_heading2(doc, compare_heading)
    add_normal(doc, _t(data, "Compare current server specification with Oracle Database 19c minimum requirements:"))

    compare_table = doc.add_table(rows=1, cols=3)
    compare_table.style = "Table Grid"
    compare_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    compare_table.autofit = False
    set_table_borders(compare_table)
    set_fixed_table_grid(compare_table, [Inches(1.10), Inches(2.35), Inches(3.75)])

    compare_headers = ["Requirement", "Minimum\nRequirement", "Current Server Specification"]
    for index, header in enumerate(compare_headers):
        cell = compare_table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_text(cell, header, bold=True, size=7.5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=8, bottom=8, left=8, right=8)

    disk_rows = []
    for disk in os_info.get("disk_info", []):
        disk_rows.append(
            {
                "Filesystem": disk.get("filesystem", ""),
                "Size": disk.get("size", ""),
                "Used": disk.get("used", ""),
                "Avail": disk.get("avail", ""),
                "Use%": disk.get("use_pct", ""),
                "Mounted\non": disk.get("mounted_on", ""),
            }
        )

    kernel_sem = os_info.get("kernel_params", {}).get("kernel.sem") or os_info.get("semaphores", "")
    current_os_text = "\n".join(
        [
            item
            for item in [
                os_info.get("os_version", ""),
                f"Linux version {os_info.get('kernel_version', '')}" if os_info.get("kernel_version") else "",
                os_info.get("architecture", ""),
            ]
            if item
        ]
    )

    comparison_rows = [
        {
            "requirement": "Linux OS:",
            "minimum": "Red Hat Enterprise Linux 8: 4.18.0-80.el8.x86_64 or later",
            "current": current_os_text,
        },
        {
            "requirement": "Linux Disk Space:",
            "minimum": "/tmp At least 1 GB\n/u01 At least 7.2 GB",
            "nested": {
                "headers": ["Filesystem", "Size", "Used", "Avail", "Use%", "Mounted\non"],
                "rows": disk_rows,
                "widths": [
                    Inches(0.78),
                    Inches(0.42),
                    Inches(0.42),
                    Inches(0.42),
                    Inches(0.42),
                    Inches(1.05),
                ],
            },
        },
        {
            "requirement": "Linux RAM:",
            "minimum": "At least 1 GB",
            "current": f"Memory size: {os_info.get('total_ram_mb', '')} MB",
        },
        {
            "requirement": "Linux Tmp:",
            "minimum": "/tmp At least 1 GB",
            "current": f"Tmp size: {os_info.get('tmp_size', '')}",
        },
        {
            "requirement": "Linux JDK & JRE:",
            "minimum": "JDK 8 / JRE 8",
            "current": os_info.get("jdk_version", ""),
        },
        {
            "requirement": "Linux kernel\nsetting:",
            "minimum": "250 32000 100 128",
            "current": kernel_sem,
        },
    ]

    for row in comparison_rows:
        cells = compare_table.add_row().cells
        set_cell_text(cells[0], row["requirement"], bold=True, size=7.5)
        set_cell_text(cells[1], row["minimum"], size=7.5)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=5, bottom=5, left=5, right=5)
        if row.get("nested"):
            add_nested_table(
                cells[2],
                row["nested"]["headers"],
                row["nested"]["rows"],
                header_fill=LIGHT_GRAY,
                font_size=5.6,
                column_widths=row["nested"].get("widths"),
            )
        else:
            set_cell_text(cells[2], row.get("current", ""), size=7.5)

    # ----- 2.3 User's environment -----
    add_heading2(doc, _t(data, "2.3 User's environment"))
    add_normal(doc, _t(data, "Oracle user environment configuration:"))

    env_rows = [
        {"Item": "Oracle Owner", "Value": os_info.get("oracle_owner", "")},
        {"Item": "Oracle Home", "Value": os_info.get("oracle_home", "")},
        {"Item": "Oracle SID", "Value": os_info.get("oracle_sid", "")},
        {"Item": "Oracle Base", "Value": os_info.get("oracle_base", "")},
        {"Item": "Shell", "Value": os_info.get("shell", "")},
        {"Item": "Groups", "Value": os_info.get("oracle_groups", "")},
        {"Item": "Semaphores", "Value": os_info.get("semaphores", "")},
    ]
    add_small_table(doc, ["Item", "Value"], env_rows)

    # Kernel parameters sub-table
    kernel_params = os_info.get("kernel_params", {})
    if kernel_params:
        add_heading2(doc, _t(data, "2.3.1 Kernel Parameters"))
        kp_rows = [{"Parameter": k, "Value": v} for k, v in kernel_params.items()]
        add_small_table(doc, ["Parameter", "Value"], kp_rows)

    # Disk space sub-table
    disk_info = os_info.get("disk_info", [])
    if disk_info:
        add_heading2(doc, _t(data, "2.3.2 Disk Space"))
        disk_rows = []
        for d in disk_info:
            disk_rows.append({
                "Filesystem": d.get("filesystem", ""),
                "Size": d.get("size", ""),
                "Used": d.get("used", ""),
                "Avail": d.get("avail", ""),
                "Use%": d.get("use_pct", ""),
                "Mounted on": d.get("mounted_on", ""),
            })

        def highlight_disk(row, h, value):
            if h == "Use%":
                try:
                    pct = int(str(value).replace("%", ""))
                    return pct >= 85
                except (ValueError, TypeError):
                    pass
            return False

        add_small_table(doc, ["Filesystem", "Size", "Used", "Avail", "Use%", "Mounted on"],
                        disk_rows, highlight_rule=highlight_disk)

    # User limits sub-table
    user_limits = os_info.get("user_limits", {})
    if user_limits:
        add_heading2(doc, _t(data, "2.3.3 User Limits (ulimit)"))
        ul_rows = [{"Limit": k, "Value": v} for k, v in user_limits.items()]
        add_small_table(doc, ["Limit", "Value"], ul_rows)

    additional_nodes = []
    primary_hostname = os_info.get("hostname", "")
    for node in visible_nodes:
        node_os = node.get("os_info", {})
        if node_os.get("hostname") and node_os.get("hostname") != primary_hostname:
            additional_nodes.append(node)

    if additional_nodes:
        add_heading2(doc, _t(data, "2.4 Additional server data from bundle ZIP"))
        for node in additional_nodes:
            node_os = node.get("os_info", {})
            add_normal(doc, f"{_t(data, 'Source')}: {node.get('source', '')}")
            add_small_table(
                doc,
                ["Item", "Detail"],
                [
                    {"Item": "Hostname", "Detail": node_os.get("hostname", "")},
                    {"Item": "IP Address", "Detail": node_os.get("ip_address", "")},
                    {"Item": "OS Version", "Detail": node_os.get("os_version", "")},
                    {"Item": "Kernel Version", "Detail": node_os.get("kernel_version", "")},
                    {"Item": "Total RAM (GB)", "Detail": node_os.get("total_ram_gb", "")},
                    {"Item": "Oracle Home", "Detail": node_os.get("oracle_home", "")},
                    {"Item": "Oracle SID", "Detail": node_os.get("oracle_sid", "")},
                    {"Item": "JDK/JRE Version", "Detail": node_os.get("jdk_version", "")},
                    {"Item": "/tmp Total Size", "Detail": node_os.get("tmp_size", "")},
                    {"Item": "/tmp Available", "Detail": node_os.get("tmp_avail", "")},
                ],
            )

    doc.add_page_break()


# =========================================================
# Section 3: System Checklist
# =========================================================

def add_system_checklist(doc, data):
    os_info = data.get("os_info", {})

    add_heading1(doc, _t(data, "3. System Checklist"))

    # ----- 3.1 Hardware configuration -----
    add_heading2(doc, _t(data, "3.1 Hardware configuration"))
    add_normal(doc, _t(data, "Hardware and system configuration details:"))

    hw_rows = [
        {"Item": "Machine Name (Hostname)", "Detail": os_info.get("hostname", "")},
        {"Item": "Platform / OS", "Detail": os_info.get("os_version", "")},
        {"Item": "Architecture", "Detail": os_info.get("architecture", "")},
        {"Item": "Kernel Version", "Detail": os_info.get("kernel_version", "")},
        {"Item": "CPU Type", "Detail": os_info.get("cpu_model", "")},
        {"Item": "Number of CPU(s)", "Detail": os_info.get("cpu_count", "")},
        {"Item": "Socket(s)", "Detail": os_info.get("sockets", "")},
        {"Item": "Core(s) per socket", "Detail": os_info.get("cores_per_socket", "")},
        {"Item": "CPU Clock Frequency (MHz)", "Detail": os_info.get("cpu_mhz", "")},
        {"Item": "CPU Cache Size", "Detail": os_info.get("cache_size", "")},
        {"Item": "Physical Memory (RAM)", "Detail": f"{os_info.get('total_ram_gb', '')} GB ({os_info.get('total_ram_mb', '')} MB)" if os_info.get("total_ram_gb") else ""},
        {"Item": "Swap Total", "Detail": f"{os_info.get('swap_total_mb', '')} MB" if os_info.get("swap_total_mb") else ""},
        {"Item": "Swap Used", "Detail": f"{os_info.get('swap_used_mb', '')} MB" if os_info.get("swap_used_mb") else ""},
    ]
    add_small_table(doc, ["Item", "Detail"], hw_rows)

    # ----- 3.2 Network configuration -----
    add_heading2(doc, _t(data, "3.2 Network configuration"))
    add_normal(doc, _t(data, "Network interface and hosts configuration:"))

    net_rows = [
        {"Item": "IP Address", "Detail": os_info.get("ip_address", "")},
        {"Item": "Subnet Mask", "Detail": os_info.get("subnet_mask", "")},
        {"Item": "Hostname", "Detail": os_info.get("hostname", "")},
    ]
    add_small_table(doc, ["Item", "Detail"], net_rows)

    # Hosts file
    hosts_entries = os_info.get("hosts_entries", [])
    if hosts_entries:
        add_heading2(doc, _t(data, "3.2.1 Hosts File (/etc/hosts)"))
        hosts_rows = [{"IP Address": h.get("ip", ""), "Hostname(s)": h.get("hostnames", "")} for h in hosts_entries]
        add_small_table(doc, ["IP Address", "Hostname(s)"], hosts_rows)

    # ----- 3.3 Crontab information -----
    add_heading2(doc, _t(data, "3.3 Crontab information"))

    crontab_entries = os_info.get("crontab_entries", [])
    if crontab_entries:
        add_normal(doc, f"Total crontab entries: {len(crontab_entries)} (Active: {sum(1 for c in crontab_entries if c.get('active'))}, Inactive: {sum(1 for c in crontab_entries if not c.get('active'))})")

        cron_rows = []
        for c in crontab_entries:
            cron_rows.append({
                "Schedule": c.get("schedule", ""),
                "Command": c.get("command", ""),
                "Status": "Active" if c.get("active") else "Inactive",
            })

        def highlight_cron(row, h, value):
            return row.get("Status") == "Inactive"

        add_small_table(doc, ["Schedule", "Command", "Status"], cron_rows, highlight_rule=highlight_cron)
    else:
        add_normal(doc, _t(data, "No crontab entries found."))

    bundle_nodes = data.get("bundle_nodes", [])
    primary_hostname = os_info.get("hostname", "")
    additional_nodes = [
        node for node in bundle_nodes
        if node.get("os_info", {}).get("hostname")
        and node.get("os_info", {}).get("hostname") != primary_hostname
    ]
    if additional_nodes:
        add_heading2(doc, _t(data, "3.4 Additional server checklist data from bundle ZIP"))
        for node in additional_nodes:
            node_os = node.get("os_info", {})
            add_normal(doc, f"{_t(data, 'Source')}: {node.get('source', '')}")
            add_small_table(
                doc,
                ["Item", "Detail"],
                [
                    {"Item": "Machine Name (Hostname)", "Detail": node_os.get("hostname", "")},
                    {"Item": "Platform / OS", "Detail": node_os.get("os_version", "")},
                    {"Item": "Kernel Version", "Detail": node_os.get("kernel_version", "")},
                    {"Item": "CPU Type", "Detail": node_os.get("cpu_model", "")},
                    {"Item": "Number of CPU(s)", "Detail": node_os.get("cpu_count", "")},
                    {"Item": "Physical Memory (RAM)", "Detail": f"{node_os.get('total_ram_gb', '')} GB ({node_os.get('total_ram_mb', '')} MB)" if node_os.get("total_ram_gb") else ""},
                    {"Item": "IP Address", "Detail": node_os.get("ip_address", "")},
                    {"Item": "Subnet Mask", "Detail": node_os.get("subnet_mask", "")},
                    {"Item": "Crontab entries", "Detail": str(len(node_os.get("crontab_entries", [])))},
                ],
            )

    doc.add_page_break()


# =========================================================
# Database Information
# =========================================================

def add_database_information(doc, data):
    add_heading1(doc, _t(data, "4. Database Information"))

    add_heading2(doc, _t(data, "4.1 Database Configuration."))

    db_name = data.get("instance_name") or data.get("db_parameters", {}).get("db_name", "")
    archive_enabled = "Yes" if data.get("archive_mode") == "ARCHIVELOG" else "No"

    rows = [
        {"What?": "Instance name", "Values?": db_name},
        {"What?": "RDBMS Version/Release?", "Values?": data.get("rdbms_version", "")},
        {"What?": "Number of datafiles?", "Values?": str(len(data.get("datafiles", [])))},
        {"What?": "Number of tempfiles?", "Values?": str(len(data.get("tempfiles", [])))},
        {"What?": "Is Archiving enabled?", "Values?": archive_enabled},
        {"What?": "Disk Space(DataFile + TempFile + RedoLogFile)", "Values?": data.get("disk_space", "")},
        {"What?": "Optimizer mode", "Values?": data.get("optimizer_mode", "")},
    ]

    add_small_table(doc, ["What?", "Values?"], rows)

    add_heading2(doc, _t(data, "4.2 Database Parameter"))

    param_rows = []
    for k, v in data.get("db_parameters", {}).items():
        param_rows.append({"Name": k, "Values": v})

    def highlight_param(row, h, value):
        return row.get("Name", "").lower() in ["audit_trail"] or row.get("Name", "").lower() == "db_file_multiblock_read_count"

    add_small_table(doc, ["Name", "Values"], param_rows, highlight_rule=highlight_param)

    add_heading2(doc, _t(data, "4.3 Major Security Initialization Parameters"))
    sec_params = []
    for k in ["O7_DICTIONARY_ACCESSIBILITY", "audit_trail", "remote_login_passwordfile", "remote_os_authent"]:
        sec_params.append({"Name": k, "Values": data.get("db_parameters", {}).get(k, "")})

    add_small_table(doc, ["Name", "Values"], sec_params, highlight_rule=highlight_param)

    add_normal(doc, _t(data, "Additional Suggestion:"), bold=True)
    add_normal(
        doc,
        "- audit_trail parameter should be reviewed. If audit data is stored in database tablespace, monitor usage and purge data according to retention policy.",
    )

    add_heading2(doc, _t(data, "4.4 Database Files"))
    df_rows = []
    for r in data.get("datafiles", []):
        df_rows.append({
            "Tbs Name": r.get("tbs_name", ""),
            "File Name": r.get("file_name", ""),
            "Size(MB)": r.get("size_mb", ""),
            "Max(MB)": r.get("max_mb", ""),
            "Aut": r.get("auto_ext", ""),
            "Inc.(MB)": r.get("inc_mb", ""),
        })

    def highlight_datafile(row, h, value):
        return row.get("Tbs Name", "") == "USERS" or str(row.get("Inc.(MB)", "")).strip() in ["1", "1.25"]

    add_small_table(doc, ["Tbs Name", "File Name", "Size(MB)", "Max(MB)", "Aut", "Inc.(MB)"], df_rows, highlight_rule=highlight_datafile)

    # Additional Suggestion for low increment datafiles
    low_inc_files = []
    for r in data.get("datafiles", []):
        try:
            inc = float(r.get("inc_mb", "0"))
            if 0 < inc < 5:
                low_inc_files.append((r.get("tbs_name", ""), inc))
        except (ValueError, TypeError):
            pass
    if low_inc_files:
        add_normal(doc, "Additional Suggestion:", bold=True)
        tbs_names = list(set(t[0] for t in low_inc_files))
        for tbs in tbs_names:
            inc_val = [v for t, v in low_inc_files if t == tbs][0]
            add_normal(doc, f"- Refer database file table, the data file of {tbs} tablespace is auto extend by increase {inc_val} MB, it's frequently auto extend size of data file. We recommend you configure auto extend by increase to 50 MB for reduce time to amount of extend size of data file.")

    add_heading2(doc, _t(data, "4.5 Temporary Files"))
    tmp_rows = []
    for r in data.get("tempfiles", []):
        tmp_rows.append({
            "Tbs Name": r.get("tbs_name", ""),
            "File Name": r.get("file_name", ""),
            "Size(MB)": r.get("size_mb", ""),
            "Max(MB)": r.get("max_mb", ""),
            "Aut": r.get("auto_ext", ""),
            "Inc.(MB)": r.get("inc_mb", ""),
        })
    add_small_table(doc, ["Tbs Name", "File Name", "Size(MB)", "Max(MB)", "Aut", "Inc.(MB)"], tmp_rows)

    add_heading2(doc, _t(data, "4.6 Redo Log File"))
    redo_rows = []
    for r in data.get("redo_logs", []):
        redo_rows.append({
            "Group#": r.get("group_no", ""),
            "Member": r.get("member", ""),
            "Size(MB)": r.get("size_mb", ""),
        })
    add_small_table(doc, ["Group#", "Member", "Size(MB)"], redo_rows)

    add_heading2(doc, _t(data, "4.7 Control File"))
    control_rows = [{"Control File Name#": r.get("file_name", "")} for r in data.get("control_files", [])]
    add_small_table(doc, ["Control File Name#"], control_rows)

    add_heading2(doc, _t(data, "4.8 Database Patch"))
    patch_rows = []
    for r in data.get("patch_info", []):
        patch_rows.append({
            "Component": r.get("component", ""),
            "Current Version": r.get("current_version", ""),
            "Recommended version": r.get("recommended_version", ""),
        })

    def highlight_patch(row, h, value):
        current = str(row.get("Current Version", "")).strip()
        recommended = str(row.get("Recommended version", "")).strip()
        return bool(current and recommended and current != recommended)

    add_small_table(
        doc,
        ["Component", "Current Version", "Recommended version"],
        patch_rows,
        header_fill=LIGHT_GRAY,
        highlight_rule=highlight_patch,
    )

    add_normal(doc, _t(data, "Additional suggestion:"), bold=True)
    add_normal(doc, "- Recommend applying the latest patch to fix bugs and update database security.")

    doc.add_page_break()


# =========================================================
# Growth graph
# =========================================================

def create_growth_chart(growth_rows, output_dir):
    if not growth_rows:
        return None

    months = [str(r.get("month", "")) for r in growth_rows]
    used = [float(r.get("used_gb", 0) or 0) for r in growth_rows]
    allocated = [float(r.get("allocated_gb", 0) or 0) for r in growth_rows]

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    chart_path = output_dir / "database_growth_rate.png"

    plt.figure(figsize=(6.2, 3.0), facecolor="#333333")
    ax = plt.gca()
    ax.set_facecolor("#333333")

    ax.plot(months, used, marker="o", label="Total use per month(GB)", linewidth=2, color="#5B9BD5")
    ax.plot(months, allocated, marker="o", label="Current Allocate (GB)", linewidth=2, color="#ED7D31")

    ax.set_title("Database Growth Rate", color="white", fontsize=11, weight="bold")
    ax.tick_params(colors="white", labelsize=8)
    ax.grid(True, color="#555555", linewidth=0.5)

    for spine in ax.spines.values():
        spine.set_color("#666666")

    legend = ax.legend(fontsize=6, loc="upper center", ncol=2, bbox_to_anchor=(0.5, 1.0))
    legend.get_frame().set_facecolor("#333333")
    legend.get_frame().set_edgecolor("#333333")
    for text in legend.get_texts():
        text.set_color("white")

    for x, y in zip(months, used):
        ax.text(x, y, f"{y:.2f}", color="white", fontsize=7, ha="center", va="bottom")

    for x, y in zip(months, allocated):
        ax.text(x, y, f"{y:.2f}", color="white", fontsize=7, ha="center", va="bottom")

    plt.tight_layout()
    plt.savefig(chart_path, dpi=180, facecolor="#333333")
    plt.close()

    return chart_path


def add_performance_review(doc, data):
    perf = data.get("perf_stats", {})
    perf_review = data.get("perf_review", {})
    db_params = data.get("db_parameters", {})
    oracle_stats = data.get("oracle_stats", {})

    add_heading2(doc, _t(data, "5.1 Performance Review"))

    outer = doc.add_table(rows=1, cols=2)
    outer.style = "Table Grid"
    outer.alignment = WD_TABLE_ALIGNMENT.LEFT
    outer.autofit = False
    set_table_borders(outer)
    set_table_column_widths(outer, [Inches(2.3), Inches(4.25)])

    set_cell_shading(outer.rows[0].cells[0], LIGHT_GRAY)
    set_cell_shading(outer.rows[0].cells[1], LIGHT_GRAY)
    set_cell_text(outer.rows[0].cells[0], "Information Required to Tune\nMemory Allocation", bold=True, size=8)
    set_cell_text(outer.rows[0].cells[1], "Answer", bold=True, size=8)

    def add_question_row(question_text, answer_text=None, answer_table=None):
        row_cells = outer.add_row().cells
        set_cell_text(row_cells[0], question_text, size=8)
        if answer_table is not None:
            add_nested_table(row_cells[1], answer_table["headers"], answer_table["rows"], header_fill=LIGHT_GRAY, font_size=7)
        else:
            set_cell_text(row_cells[1], answer_text or "", size=8)
        return row_cells

    library_rows = [
        {"NameSpace": row.get("namespace", ""), "GetHitRatio": row.get("gethitratio", "")}
        for row in perf_review.get("library_cache_namespace", [])
    ]
    add_question_row(
        "1. What is the gethitratio of the\nlibrarycache?",
        answer_table={
            "headers": ["NameSpace", "GetHitRatio"],
            "rows": library_rows or [{"NameSpace": "", "GetHitRatio": perf.get("LIBRARY_CACHE_HITRATIO", "")}],
        },
    )

    pin_reload = perf_review.get("pin_reload", {})
    add_question_row(
        '2. What is the PIN / RELOAD ratio\nwithin the librarycache;\n\nSelect sum(pins) "Executions",\nsum(reloads) "Cache Misses",\nsum(reloads)/sum(pins) from\nv$librarycache;\n\nNote: Reload should ideally be\nZERO\nNever more than 1% of the PINS.',
        answer_table={
            "headers": ["Execution", "Cache Misses", "Sum"],
            "rows": [
                {
                    "Execution": pin_reload.get("executions", ""),
                    "Cache Misses": pin_reload.get("cache_misses", ""),
                    "Sum": pin_reload.get("ratio", ""),
                }
            ],
        },
    )

    def add_simple_pair(label, value):
        row_cells = outer.add_row().cells
        set_cell_text(row_cells[0], label, size=8)
        set_cell_text(row_cells[1], value, size=8)

    add_simple_pair("If reloads to pin ratio is >1% -\nINCREASE the shared_pool_size.", "")
    add_simple_pair("3. Data Dictionary Cache Miss\nRatio\nKeep this below 5%", f"Dictionary Cache Hit Ratio: {perf.get('DICT_CACHE_HITRATIO', '')}")
    add_simple_pair("4. shared_pool_size =?", db_params.get("shared_pool_size", ""))
    add_simple_pair("5. shared_pool_reserved_size =?", db_params.get("shared_pool_reserved_size", ""))

    shared_pool_stats = perf_review.get("shared_pool_stats", {})
    shared_pool_lines = [
        f"Free space: {shared_pool_stats.get('free_space', '')}",
        f"Average Free Size: {shared_pool_stats.get('avg_free_size', '')}",
        f"Max Free Size: {shared_pool_stats.get('max_free_size', '')}",
        f"Used Space: {shared_pool_stats.get('used_space', '')}",
        f"Average Used Space: {shared_pool_stats.get('avg_used_size', '')}",
    ]
    add_simple_pair(
        "6. What are the\nSHARED_POOL_RESERVED\nstatistics?\nSelect * from\nv$shared_pool_reserved?",
        "\n".join([line for line in shared_pool_lines if not line.endswith(": ")]),
    )
    add_simple_pair("7. Redo log space request", perf.get("REDO_LOG_SPACE_REQUEST", ""))
    add_simple_pair("8. DB Block Buffer Cache Hit\nRatio?", perf.get("BUFFER_CACHE_HIT_RATIO", ""))
    add_simple_pair("9. Latch Hit Ratio?", perf.get("LATCH_HIT_RATIO", ""))
    add_simple_pair("10. Disk Sort Ratio?", perf.get("DISK_SORT_RATIO", ""))
    add_simple_pair("11. Rollback Segment Waits?", perf.get("ROLLBACK_SEGMENT_WAITS", ""))
    add_simple_pair("12. Dispatcher Workload?", perf.get("DISPATCHER_WORKLOAD", ""))
    add_simple_pair("13. PGA cache hit percentage", perf.get("PGA_CACHE_HIT_PCT", oracle_stats.get("PGA_CACHE_HIT_PCT", "")))
    add_simple_pair("The UNDO Tablespace", db_params.get("undo_tablespace", ""))

    undo_rows = [
        {"Amount": row.get("amount", ""), "Segment Type": row.get("segment_type", ""), "Size(MB)": row.get("size_mb", "")}
        for row in perf_review.get("undo_segments", [])
    ]
    add_question_row(
        "Number and size of Undo\nSegments?",
        answer_table={
            "headers": ["Amount", "Segment Type", "Size(MB)"],
            "rows": undo_rows or [{"Amount": "", "Segment Type": "", "Size(MB)": ""}],
        },
    )

    add_simple_pair("Information Required to Tune\nLogging and Archiving", "")
    add_simple_pair("At least 3 redo log groups?", "Yes" if len(data.get("redo_logs", [])) >= 3 else "No")
    add_simple_pair("Are using Archive Mode", "Yes" if data.get("archive_mode", "") == "ARCHIVELOG" else "No")

    archive_format = str(db_params.get("log_archive_format", ""))
    has_sequence_number = "%s" in archive_format.lower() or "%S" in archive_format
    add_simple_pair("Archive log names include\nsequence number?", "Yes" if has_sequence_number else "No")

    if data.get("archive_mode", "") != "ARCHIVELOG":
        add_normal(doc, _t(data, "Additional Suggestion:"), bold=True)
        add_normal(doc, "- An archived redo log file is a copy of one of the filled members of a redo log group. This database is in NOARCHIVELOG indicates you disable the archiving of the redo log. If a media failure occurs while the database is in NOARCHIVELOG mode, you can only restore the database to the point of the most recent full database backup. You cannot recover transactions subsequent to that backup.")
        add_normal(doc, "- We recommend you run a database in ARCHIVELOG mode indicates you enable the archiving of the redo log. A database backup, together with online and archived redo log files, guarantees that you can recover all committed transactions in the event of an operating system or disk failure. When your database is in archivelog mode, you have to have more disk space to store archive log file and it will decrease your database performance because Archiver Processes (ARCn) copy redo log files to a storage device.")


def add_database_growth_rate(doc, data, output_dir):
    add_heading1(doc, _t(data, "5. RDBMS Performance"))
    add_performance_review(doc, data)
    add_heading2(doc, _t(data, "5.2 Database Growth Rate"))

    growth_rows = data.get("growth_rows_raw", [])
    db_name = data.get("instance_name") or data.get("db_parameters", {}).get("db_name", "")

    if not growth_rows:
        add_normal(doc, "No database growth data found.")
        return

    display_rows = growth_rows[-4:] if len(growth_rows) > 4 else growth_rows
    current = display_rows[-1]
    current_used = current.get("used_gb", "")
    current_alloc = current.get("allocated_gb", "")
    current_pct = current.get("pct_used", "")

    add_normal(doc, f"Current data information about space usage of database {db_name} as follows")

    add_small_table(doc, ["Name", "Value"], [
        {"Name": "Database Name", "Value": db_name},
        {"Name": "Current Allocated (GB)", "Value": current_alloc},
        {"Name": "Current Used data (GB)", "Value": current_used},
        {"Name": "Percent Used", "Value": current_pct},
    ], header_fill=HEADER_BLUE)

    chart_path = create_growth_chart(display_rows, output_dir)
    if chart_path:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(chart_path), width=Inches(6.2))

    add_normal(doc, "The result of total database usage size in this quarter as the below table.")

    table_rows = []
    for r in display_rows:
        table_rows.append({
            "Month": r.get("month", ""),
            "Current Used (GB)": r.get("used_gb", ""),
            "Current Allocate (GB)": r.get("allocated_gb", ""),
            "% Used": r.get("pct_used", ""),
        })

    add_small_table(doc, ["Month", "Current Used (GB)", "Current Allocate (GB)", "% Used"], table_rows, header_fill=HEADER_BLUE)

    add_normal(
        doc,
        f"The current size of the database is {current_used} GB. ",
    )

    # Dynamic growth rate analysis
    if len(display_rows) >= 2:
        first = display_rows[0]
        last = display_rows[-1]
        used_diff = last.get("used_gb", 0) - first.get("used_gb", 0)
        n_months = max(len(display_rows) - 1, 1)
        avg_growth_per_month = round(used_diff / n_months, 2)

        alloc_diff = last.get("allocated_gb", 0) - first.get("allocated_gb", 0)
        avg_alloc_per_month = round(alloc_diff / n_months, 2)

        if avg_growth_per_month > 0:
            add_normal(
                doc,
                f"We compare data growth rate between current allocated data which is {current_alloc} GB "
                f"and the allocated size increasing around {avg_alloc_per_month} GB per month. "
                f"The database growth rate is around {avg_growth_per_month} GB per month. "
                f"As a result of database size, the current allocated size of the database is sufficient to support database growth rate in this current situation.",
            )
        else:
            add_normal(
                doc,
                f"The current allocated size is {current_alloc} GB. "
                f"The database growth rate decreased by around {abs(avg_growth_per_month)} GB per month. "
                f"As a result of database size, the current allocated size of the database is sufficient to support database growth rate in this current situation.",
            )
    else:
        add_normal(
            doc,
            f"The current allocated size is {current_alloc} GB. "
            f"As a result of database size, the current allocated size of the database is sufficient to support database growth rate in this current situation.",
        )

    doc.add_page_break()


# =========================================================
# Section 5.3: Performance Analysis
# =========================================================

def _fmt_number(value, decimals=0):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if decimals:
        return f"{number:,.{decimals}f}"
    return f"{number:,.0f}"


def _fmt_plain_number(value, decimals=0):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if decimals:
        return f"{number:.{decimals}f}".rstrip("0").rstrip(".")
    return f"{number:.0f}"


def _current_index(rows):
    for index, row in enumerate(rows):
        if row.get("is_current"):
            return index
    factors = []
    for index, row in enumerate(rows):
        factor = row.get("pga_target_factor", row.get("sga_size_factor", None))
        try:
            factors.append((abs(float(factor) - 1.0), index))
        except (TypeError, ValueError):
            pass
    return min(factors)[1] if factors else 0


def _advisory_window(rows, before=1, after=3):
    if not rows:
        return []
    current = _current_index(rows)
    start = max(0, current - before)
    end = min(len(rows), current + after + 1)
    if end - start < before + after + 1:
        start = max(0, end - (before + after + 1))
    return rows[start:end]


def _dark_advisory_chart(rows, x_labels, y_values, title, x_label, y_label, legend_label, output_path):
    fig, ax = plt.subplots(figsize=(7.1, 3.5), dpi=150, facecolor="#2F2F2F")
    ax.set_facecolor("#3B3B3B")

    ax.plot(x_labels, y_values, marker="o", color="#ED7D31", linewidth=2.2, markersize=4, label=legend_label)
    ax.set_title(title, color="#EDEDED", fontsize=13, fontweight="bold", pad=12)
    ax.set_xlabel(x_label, color="#D9D9D9", fontsize=8, fontweight="bold")
    ax.set_ylabel(y_label, color="#D9D9D9", fontsize=8, fontweight="bold")
    ax.tick_params(colors="#D9D9D9", labelsize=7)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value:,.0f}"))
    ax.grid(True, axis="y", color="#5A5A5A", linewidth=0.6)

    for spine in ax.spines.values():
        spine.set_visible(False)

    for x, y in zip(x_labels, y_values):
        ax.text(x, y, _fmt_number(y), color="#CFCFCF", fontsize=7, ha="center", va="bottom")

    current = _current_index(rows)
    if 0 <= current < len(rows):
        current_x = x_labels[current]
        current_y = y_values[current]
        ax.scatter([current_x], [current_y], s=90, facecolors="none", edgecolors="yellow", linewidths=2.2, zorder=4)
        ax.annotate(
            "Current",
            xy=(current_x, current_y),
            xytext=(-55, -25),
            textcoords="offset points",
            bbox={"boxstyle": "square,pad=0.35", "fc": "yellow", "ec": "black", "lw": 0.8},
            color="black",
            fontsize=8,
            arrowprops={"arrowstyle": "-", "color": "yellow", "lw": 1.2},
        )

    legend = ax.legend(loc="lower center", bbox_to_anchor=(0.5, -0.28), fontsize=7, frameon=False)
    for text in legend.get_texts():
        text.set_color("#D9D9D9")

    fig.tight_layout(pad=2.0)
    fig.savefig(output_path, facecolor="#2F2F2F")
    plt.close(fig)


def create_pga_advisory_chart(pga_rows, output_dir):
    if not pga_rows:
        return None

    rows = _advisory_window(pga_rows)
    labels = []
    values = []
    for row in rows:
        size_mb = int(float(row.get("pga_size_bytes", 0) or 0) / (1024 * 1024))
        label = _fmt_number(size_mb)
        if row.get("is_current"):
            label = f"{label} (Current)"
        labels.append(label)
        values.append(float(row.get("estd_extra_mb_rw", 0) or 0))

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    chart_path = output_dir / "database_pga_advisory.png"

    _dark_advisory_chart(
        rows,
        labels,
        values,
        title="PGA Memory Advisory",
        x_label="PGA TARGET EST (MB)",
        y_label="ESTD EXTRA READ/WRITTEN TO DISK",
        legend_label="Estd Extra W/A MB Read/Written to Disk",
        output_path=chart_path,
    )
    
    return chart_path
    
def create_sga_advisory_chart(sga_rows, output_dir):
    if not sga_rows:
        return None

    rows = _advisory_window(sga_rows)
    labels = []
    values = []
    for row in rows:
        size_mb = int(float(row.get("sga_size_mb", 0) or 0))
        label = _fmt_number(size_mb)
        if row.get("is_current"):
            label = f"{label} (Current)"
        labels.append(label)
        values.append(float(row.get("estd_physical_reads", 0) or 0))

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    chart_path = output_dir / "database_sga_advisory.png"

    _dark_advisory_chart(
        rows,
        labels,
        values,
        title="SGA Memory Advisory",
        x_label="SGA TARGET SIZE (MB)",
        y_label="EST PHYSICAL READS",
        legend_label="Est Physical Reads",
        output_path=chart_path,
    )
    
    return chart_path


def _sample_tick_labels(samples, max_ticks=10):
    if not samples:
        return [], []

    step = max(1, len(samples) // max_ticks)
    positions = list(range(0, len(samples), step))
    if positions[-1] != len(samples) - 1:
        positions.append(len(samples) - 1)

    labels = [samples[pos].get("day", "") for pos in positions]
    return positions, labels


def _create_oswatcher_chart(samples, value_key, chart_title, output_path, y_limit=None):
    if not samples:
        return None

    output_path = Path(output_path)
    values = [float(row.get(value_key, 0) or 0) for row in samples]
    x_values = list(range(len(values)))
    tick_positions, tick_labels = _sample_tick_labels(samples)

    fig, ax = plt.subplots(figsize=(8.2, 2.0), dpi=140)
    fig.patch.set_facecolor("#C9C7EF")
    ax.set_facecolor("#BFBFBF")
    ax.plot(x_values, values, color="#002BFF", linewidth=0.8)

    ax.set_title(chart_title, fontsize=7, fontweight="bold", pad=3)
    ax.grid(True, color="#777777", linewidth=0.35, alpha=0.65)
    ax.tick_params(axis="both", labelsize=5, pad=1)
    ax.set_xlim(0, max(len(values) - 1, 1))
    if y_limit:
        ax.set_ylim(0, y_limit)

    ax.set_xticks(tick_positions)
    ax.set_xticklabels(tick_labels)
    for spine in ax.spines.values():
        spine.set_color("#666666")
        spine.set_linewidth(0.6)

    fig.text(0.02, 0.02, "Host=MotifPRDDBN", fontsize=5, color="#333333")
    plt.tight_layout(pad=0.8)
    plt.savefig(output_path, facecolor=fig.get_facecolor())
    plt.close(fig)
    return output_path


def _create_oswatcher_multi_chart(samples, series, chart_title, output_path, y_limit=None):
    if not samples:
        return None

    output_path = Path(output_path)
    x_values = list(range(len(samples)))
    tick_positions, tick_labels = _sample_tick_labels(samples)

    fig, ax = plt.subplots(figsize=(8.2, 2.0), dpi=140)
    fig.patch.set_facecolor("#C9C7EF")
    ax.set_facecolor("#BFBFBF")

    for key, label, color in series:
        values = [float(row.get(key, 0) or 0) for row in samples]
        ax.plot(x_values, values, color=color, linewidth=0.8, label=label)

    ax.set_title(chart_title, fontsize=7, fontweight="bold", pad=3)
    ax.grid(True, color="#777777", linewidth=0.35, alpha=0.65)
    ax.tick_params(axis="both", labelsize=5, pad=1)
    ax.set_xlim(0, max(len(samples) - 1, 1))
    if y_limit:
        ax.set_ylim(0, y_limit)
    ax.set_xticks(tick_positions)
    ax.set_xticklabels(tick_labels)
    if len(series) > 1:
        ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.18), ncol=min(len(series), 4), fontsize=5, frameon=False)
    for spine in ax.spines.values():
        spine.set_color("#666666")
        spine.set_linewidth(0.6)

    fig.text(0.02, 0.02, "Host=MotifPRDDBN", fontsize=5, color="#333333")
    plt.tight_layout(pad=0.8)
    plt.savefig(output_path, facecolor=fig.get_facecolor())
    plt.close(fig)
    return output_path


def _max_disk_samples_by_time(samples):
    by_key = {}
    for row in samples:
        key = (row.get("day", ""), row.get("time", ""))
        current = by_key.get(key)
        if current is None or float(row.get("svctm", 0) or 0) > float(current.get("svctm", 0) or 0):
            by_key[key] = row
    return list(by_key.values())


def _rate_percent(samples, key, threshold):
    if not samples:
        return 0.0
    count = sum(1 for row in samples if float(row.get(key, 0) or 0) > threshold)
    return (count / len(samples)) * 100.0


def create_os_perf_charts(os_perf_data, output_dir):
    """Generate CPU, Memory, and Disk line charts using matplotlib"""
    if not os_perf_data or not os_perf_data.get("found"):
        return None, None, None
        
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    cpu_path = output_dir / "os_cpu_perf.png"
    mem_path = output_dir / "os_mem_perf.png"
    disk_path = output_dir / "os_disk_perf.png"
    runq_path = output_dir / "os_runq_perf.png"
    cpu_all_path = output_dir / "os_cpu_util_all.png"
    cpu_sys_path = output_dir / "os_cpu_util_system.png"
    cpu_usr_path = output_dir / "os_cpu_util_user.png"
    mem_free_path = output_dir / "os_mem_free.png"
    swap_used_path = output_dir / "os_swap_used.png"
    page_in_path = output_dir / "os_page_in.png"
    page_out_path = output_dir / "os_page_out.png"
    cpu_wait_io_path = output_dir / "os_cpu_wait_io.png"
    disk_service_path = output_dir / "os_disk_service_time.png"

    runq_samples = os_perf_data.get("runq_samples", [])
    cpu_samples = os_perf_data.get("cpu_samples", [])
    mem_free_samples = os_perf_data.get("mem_free_samples", [])
    swap_samples = os_perf_data.get("swap_samples", [])
    page_io_samples = os_perf_data.get("page_io_samples", [])
    disk_samples = _max_disk_samples_by_time(os_perf_data.get("disk_samples", []))
    os_perf_data["chart_paths"] = {
        "runq": str(_create_oswatcher_chart(runq_samples, "runq", "OSWatcher  CPU: Run Queue", runq_path) or ""),
        "cpu_all": str(_create_oswatcher_chart(cpu_samples, "util", "OSWatcher  CPU Utilization", cpu_all_path, y_limit=100) or ""),
        "cpu_sys": str(_create_oswatcher_chart(cpu_samples, "sys", "OSWatcher  CPU Utilization: System", cpu_sys_path, y_limit=100) or ""),
        "cpu_usr": str(_create_oswatcher_chart(cpu_samples, "usr", "OSWatcher  CPU Utilization: User", cpu_usr_path, y_limit=100) or ""),
        "mem_free": str(_create_oswatcher_multi_chart(mem_free_samples, [("free_gb", "Memory Free (GB)", "#005A7F"), ("inactive_gb", "Inactive Cache (GB)", "#F0782E")], "OSWatcher Memory Free (Gbytes)", mem_free_path) or ""),
        "swap_used": str(_create_oswatcher_chart(swap_samples, "used_gb", "OSWatcher  Memory: Swap Used (G Bytes)", swap_used_path) or ""),
        "page_in": str(_create_oswatcher_chart(page_io_samples, "pgpgin", "OSWatcher Memory: Page In Rate (Pages Per Second)", page_in_path) or ""),
        "page_out": str(_create_oswatcher_chart(page_io_samples, "pgpgout", "OSWatcher Memory: Page Out Rate (Pages Per Second)", page_out_path) or ""),
        "cpu_wait_io": str(_create_oswatcher_chart(cpu_samples, "iowait", "OSWatcher  CPU Utilization: Wait I/O", cpu_wait_io_path, y_limit=100) or ""),
        "disk_service": str(_create_oswatcher_multi_chart(disk_samples, [("svctm", "Service Time (ms)", "#1451FF"), ("await", "I/O Wait (ms)", "#E3342F")], "OSWatcher  IO: Devices With Highest Service Time In Milliseconds", disk_service_path) or ""),
    }
    
    # 1. CPU Chart
    cpu_trend = os_perf_data.get("cpu_trend", [])
    if cpu_trend:
        days = [r["day"] for r in cpu_trend]
        usr = [r["usr"] for r in cpu_trend]
        sys = [r["sys"] for r in cpu_trend]
        iowait = [r["iowait"] for r in cpu_trend]
        
        plt.figure(figsize=(10, 5), dpi=120)
        plt.plot(days, usr, marker='o', label='%usr', color='#00478F', linewidth=2)
        plt.plot(days, sys, marker='s', label='%sys', color='#FF8C00', linewidth=2)
        plt.plot(days, iowait, marker='^', label='%iowait', color='#E63946', linewidth=2)
        
        plt.title("CPU Utilization Trend (Daily Average)", fontsize=14, fontweight="bold", pad=15)
        plt.xlabel("Day of Month", fontsize=10)
        plt.ylabel("CPU %", fontsize=10)
        plt.legend()
        plt.grid(True, linestyle="--", alpha=0.6)
        plt.xticks(rotation=45)
        plt.tight_layout()
        plt.savefig(cpu_path)
        plt.close()
    else:
        cpu_path = None
        
    # 2. Memory Chart (Page In/Out or Swap)
    mem_trend = os_perf_data.get("mem_trend", [])
    if mem_trend:
        days = [r["day"] for r in mem_trend]
        pgpgin = [r["pgpgin"] for r in mem_trend]
        pgpgout = [r["pgpgout"] for r in mem_trend]
        
        plt.figure(figsize=(10, 5), dpi=120)
        plt.plot(days, pgpgin, marker='o', label='Page In (pgpgin/s)', color='#2A9D8F', linewidth=2)
        plt.plot(days, pgpgout, marker='s', label='Page Out (pgpgout/s)', color='#E76F51', linewidth=2)
        
        plt.title("Memory Page In/Out Trend (Daily Average)", fontsize=14, fontweight="bold", pad=15)
        plt.xlabel("Day of Month", fontsize=10)
        plt.ylabel("Pages/s", fontsize=10)
        plt.legend()
        plt.grid(True, linestyle="--", alpha=0.6)
        plt.xticks(rotation=45)
        plt.tight_layout()
        plt.savefig(mem_path)
        plt.close()
    else:
        mem_path = None
        
    # 3. Disk Chart (Service Time)
    disk_trend = os_perf_data.get("disk_trend", [])
    if disk_trend:
        days = [r["day"] for r in disk_trend]
        svctm = [r["svctm"] for r in disk_trend]
        await_val = [r["await"] for r in disk_trend]
        
        plt.figure(figsize=(10, 5), dpi=120)
        plt.plot(days, svctm, marker='o', label='Service Time (svctm)', color='#6D6875', linewidth=2)
        plt.plot(days, await_val, marker='s', label='I/O Wait (await)', color='#B5838D', linewidth=2)
        
        plt.title("Disk Performance Trend (Daily Average)", fontsize=14, fontweight="bold", pad=15)
        plt.xlabel("Day of Month", fontsize=10)
        plt.ylabel("Milliseconds (ms)", fontsize=10)
        plt.legend()
        plt.grid(True, linestyle="--", alpha=0.6)
        plt.xticks(rotation=45)
        plt.tight_layout()
        plt.savefig(disk_path)
        plt.close()
    else:
        disk_path = None
        
    return cpu_path, mem_path, disk_path

def add_performance_analysis(doc, data, output_dir):
    sec53 = data.get("section53", {})
    
    add_heading2(doc, _t(data, "5.3 Performance Analysis"))
    
    # ----- 5.3.1 CPU Section -----
    add_heading3(doc, "5.3.1 CPU Section")
    add_normal(doc, "Instance Efficiency Percentages (Target 100%)", bold=True)
    inst_eff = sec53.get("instance_efficiency", {})
    if inst_eff:
        rows = [{"Metric": k, "Percentage (%)": v} for k, v in inst_eff.items()]
        add_small_table(doc, ["Metric", "Percentage (%)"], rows, header_fill=HEADER_BLUE)
        
        parse_cpu_elapsd = inst_eff.get("Parse CPU to Parse Elapsd %", "")
        try:
            parse_val = float(parse_cpu_elapsd)
            if parse_val < 50.0:
                add_normal(doc, f"Parse CPU to Parse Elapsed is low ratio ({parse_val}%), more CPU time is spending to parse SQL statements. It may mean investigation is warranted to determine why the CPU is spending this much time simply parsing SQL statements. However, if this is normal behavior of the database, it can be ignored.")
            else:
                add_normal(doc, "Parse CPU to Parse Elapsed ratio is normal.")
        except ValueError:
            pass
            
        add_normal(doc, "Other elements of instance efficiency percentages are fine.")
    else:
        add_normal(doc, "No instance efficiency data found in fast assessment logs.")
    
    # ----- 5.3.2 SQL Section -----
    add_heading3(doc, "5.3.2 SQL Section")
    
    add_normal(doc, "Top 5 Timed Foreground Events", bold=True)
    top5_events = sec53.get("top5_events", [])
    if top5_events:
        for event in top5_events:
            add_normal(doc, f"- {event}")
        if any("dblink" in ev.lower() for ev in top5_events):
            add_normal(doc, "- The SQL*Net message from dblink wait event means that your local system is waiting on the network to transfer the data across the network. It's a very normal wait event for this sort of query. We recommend checking and improving the Bandwidth network or creating a view on the remote site referencing the local tables.")
    else:
        add_normal(doc, "No Top 5 Timed Events found.")
    
    add_normal(doc, "")
    
    add_normal(doc, "Top SQL by Disk Reads:")
    sql_disk = sec53.get("top_sql_disk", [])
    if sql_disk:
        rows = [{"Reads": q.get("reads"), "SQL_ID": q.get("sql_id"), "Text": q.get("text")} for q in sql_disk]
        add_small_table(doc, ["Reads", "SQL_ID", "Text"], rows, header_fill=HEADER_BLUE)
    else:
        add_normal(doc, "No Top SQL by Disk Reads found.")
        
    add_normal(doc, "Top SQL by Buffer Gets (Memory):")
    sql_mem = sec53.get("top_sql_mem", [])
    if sql_mem:
        rows = [{"Buffer Gets": q.get("buffer_gets"), "SQL_ID": q.get("sql_id"), "Text": q.get("text")} for q in sql_mem]
        add_small_table(doc, ["Buffer Gets", "SQL_ID", "Text"], rows, header_fill=HEADER_BLUE)
    else:
        add_normal(doc, "No Top SQL by Buffer Gets found.")
        
    # ----- 5.3.3 Memory Section -----
    add_heading3(doc, "5.3.3 Memory Section")
    mem_pools = sec53.get("memory_pools", {})
    if mem_pools.get("shared_free_pct"):
        rows = [
            {"Pool Name": "Shared Pool Free %", "Value": mem_pools.get("shared_free_pct", "")},
            {"Pool Name": "Java Pool Free %", "Value": mem_pools.get("java_free_pct", "")},
            {"Pool Name": "Large Pool Free %", "Value": mem_pools.get("large_free_pct", "")},
        ]
        add_small_table(doc, ["Pool Name", "Value"], rows)
    else:
        add_normal(doc, "No detailed memory pool statistics found.")
        
    # ----- 5.3.4 Oracle Open Cursors -----
    add_heading3(doc, "5.3.4 Oracle Open Cursors")
    cursors = sec53.get("open_cursors", {})
    if cursors.get("limit"):
        rows = [
            {"Parameter": "Open Cursors Limit", "Value": cursors.get("limit", "")},
            {"Parameter": "Estimated Usage", "Value": cursors.get("usage", "")},
            {"Parameter": "Usage %", "Value": cursors.get("pct", "")},
        ]
        add_small_table(doc, ["Parameter", "Value"], rows)
        
        try:
            pct_val = float(cursors.get("pct", "0").replace("%", ""))
            if pct_val >= 95.0:
                add_normal(doc, f"As above table shows that percent usage of open_cursors and session_cached_cursors are {pct_val}%. It indicates that open_cursors and session_cached_cursors of the database are insufficient. Therefore, we recommend increasing the value of open_cursors and session_cached_cursors to 2000. It may improve these values. However, it will consume more memory in PGA. So, you should increase the PGA size too.")
            else:
                add_normal(doc, f"The percent usage of open_cursors is {pct_val}%, which is within acceptable limits.")
        except ValueError:
            pass
    else:
        add_normal(doc, "No open cursors data found.")
        
    # ----- 5.3.5 PGA Memory Advisory -----
    add_heading3(doc, "Program Global Area (PGA) Analysis")
    pga_rows = sec53.get("pga_advisory", [])
    if pga_rows:
        pga_display_rows = _advisory_window(pga_rows)
        current_pga = pga_rows[_current_index(pga_rows)]
        current_pga_mb = int(float(current_pga.get("pga_size_bytes", 0) or 0) / (1024 * 1024))
        current_pga_hit = current_pga.get("estd_pga_cache_hit_pct", 0)
        current_extra_mb = current_pga.get("estd_extra_mb_rw", 0)

        summary_table = add_small_table(
            doc,
            ["Allocate (MB)", "PGA cache Hit %"],
            [{"Allocate (MB)": _fmt_number(current_pga_mb), "PGA cache Hit %": _fmt_number(current_pga_hit, 0)}],
            header_fill=HEADER_BLUE,
        )
        summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        chart_path = create_pga_advisory_chart(pga_rows, output_dir)
        if chart_path:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(6.7))

        rows = []
        for row in pga_display_rows:
            pga_size_mb = int(float(row.get("pga_size_bytes", 0) or 0) / (1024 * 1024))
            label = _fmt_number(pga_size_mb)
            if row.get("is_current"):
                label = f"{label} (Current)"
            rows.append(
                {
                    "PGA Target\nEst (MB)": label,
                    "Estd Extra   W/A MB\nRead/Written to Disk": _fmt_number(row.get("estd_extra_mb_rw", 0)),
                    "Estd PGA\nCache Hit %": _fmt_number(row.get("estd_pga_cache_hit_pct", 0), 1),
                    "_current": row.get("is_current", False),
                }
            )

        def highlight_current(row, h, value):
            return bool(row.get("_current"))

        pga_table = add_small_table(
            doc,
            ["PGA Target\nEst (MB)", "Estd Extra   W/A MB\nRead/Written to Disk", "Estd PGA\nCache Hit %"],
            rows,
            header_fill=HEADER_BLUE,
            highlight_rule=highlight_current,
        )
        pga_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        add_normal(
            doc,
            f"From the graph PGA Target size forecast shows that current PGA Target size is sufficient. "
            f"The current PGA Target size is {_fmt_number(current_pga_mb)} MB, it has estimated Read/Written to Disk "
            f"{_fmt_number(current_extra_mb)} and PGA cache hit {_fmt_number(current_pga_hit, 0)}%.",
        )
    else:
        add_normal(doc, "No PGA Advisory data found.")
        
    # ----- 5.3.6 SGA Memory Advisory -----
    add_heading3(doc, "System Global Area (SGA) Analysis")
    sga_rows = sec53.get("sga_advisory", [])
    if sga_rows:
        sga_display_rows = _advisory_window(sga_rows)
        current_sga = sga_rows[_current_index(sga_rows)]
        current_sga_mb = int(float(current_sga.get("sga_size_mb", 0) or 0))
        current_sga_reads = int(float(current_sga.get("estd_physical_reads", 0) or 0))

        chart_path = create_sga_advisory_chart(sga_rows, output_dir)
        if chart_path:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(6.7))

        rows = []
        for row in sga_display_rows:
            sga_size_mb = int(float(row.get("sga_size_mb", 0) or 0))
            label = _fmt_number(sga_size_mb)
            if row.get("is_current"):
                label = f"{label} (Current)"
            rows.append(
                {
                    "SGA Target\nSize (MB)": label,
                    "Est Physical Reads": _fmt_number(row.get("estd_physical_reads", 0)),
                    "_current": row.get("is_current", False),
                }
            )

        def highlight_current_sga(row, h, value):
            return bool(row.get("_current"))

        sga_table = add_small_table(
            doc,
            ["SGA Target\nSize (MB)", "Est Physical Reads"],
            rows,
            header_fill=HEADER_BLUE,
            highlight_rule=highlight_current_sga,
        )
        sga_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        add_normal(
            doc,
            f"From the graph, the SGA Target size forecast shows that the current SGA Target size is sufficient. "
            f"The current SGA Target size is {_fmt_number(current_sga_mb)} MB, it has estimate physical reads "
            f"{_fmt_number(current_sga_reads)}.",
        )
    else:
        add_normal(doc, "No SGA Advisory data found.")
        
    doc.add_page_break()


# =========================================================
# Tablespace / Registry / Appendix
# =========================================================

def add_tablespace_free_space(doc, data):
    add_heading1(doc, _t(data, "6. Tablespace Free Space"))
    add_heading2(doc, _t(data, "6.1 Tablespace Free Space"))

    rows = []
    for entry in _build_tablespace_capacity_rows(data):
        rows.append(
            {
                "TABLESPACE_NAME": entry["name"],
                "Allocated (MB)": _fmt_plain_number(entry["allocated_mb"]),
                "Used (MB)": _fmt_plain_number(entry["used_mb"], 2),
                "Max (MB)": _fmt_plain_number(entry["max_mb"], 2),
                "Percentage of Free space (%)": _fmt_plain_number(entry["free_pct_of_max"], 2),
                "Free of Max (MB)": _fmt_plain_number(entry["free_of_max_mb"], 2),
                "_status": entry["status"],
            }
        )

    headers = [
        "TABLESPACE_NAME",
        "Allocated\n(MB)",
        "Used (MB)",
        "Max (MB)",
        "Percentage\nof Free\nspace (%)",
        "Free of Max\n(MB)",
    ]
    value_keys = [
        "TABLESPACE_NAME",
        "Allocated (MB)",
        "Used (MB)",
        "Max (MB)",
        "Percentage of Free space (%)",
        "Free of Max (MB)",
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_borders(table)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, HEADER_BLUE)
        set_cell_text(cell, header, bold=True, color=WHITE, size=8, align="center")

    for row in rows:
        cells = table.add_row().cells
        row_fill = YELLOW if row.get("_status") == "Critical" else WHITE
        for index, key in enumerate(value_keys):
            set_cell_shading(cells[index], row_fill)
            align = "left" if index == 0 else "right"
            set_cell_text(cells[index], row.get(key, ""), size=8, align=align)

    doc.add_page_break()


def add_default_tablespace(doc, data):
    add_heading1(doc, _t(data, "7. Default tablespace and temporary tablespace"))
    add_heading2(doc, _t(data, "7.1 Default tablespace and temporary tablespace"))

    rows = []
    for r in data.get("user_tablespaces", []):
        rows.append({
            "User Name": r.get("user_name", ""),
            "Default Tablespace": r.get("default_tbs", ""),
            "Temporary Tablespace": r.get("temp_tbs", ""),
        })

    add_small_table(doc, ["User Name", "Default Tablespace", "Temporary Tablespace"], rows)

    doc.add_page_break()


def add_database_registry(doc, data):
    add_heading1(doc, _t(data, "8. Database Registry"))
    add_heading2(doc, _t(data, "8.1 Check Database Registry"))

    rows = []
    for r in data.get("registry_list", []):
        rows.append({
            "Comp ID": r.get("comp_id", ""),
            "Version": r.get("version", ""),
            "Status": r.get("status", ""),
            "Last Modified": r.get("last_modified", ""),
        })

    add_small_table(doc, ["Comp ID", "Version", "Status", "Last Modified"], rows)

    doc.add_page_break()


def add_appendix_invalid_objects(doc, data):
    add_heading1(doc, _t(data, "APPENDIX A – Invalid Object"))
    add_heading2(doc, _t(data, "A.a Invalid Objects report."))

    rows = []
    for r in data.get("invalid_objects_list", []):
        rows.append({
            "Creator": r.get("creator", ""),
            "Object Name": r.get("object_name", ""),
            "Object Type": r.get("object_type", ""),
            "Status": r.get("status", ""),
        })

    add_small_table(doc, ["Creator", "Object Name", "Object Type", "Status"], rows[:350])

    if len(rows) > 350:
        add_normal(doc, f"Showing first 350 of {len(rows)} invalid objects.")

    add_heading2(doc, _t(data, "A.b Disabled Constraints."))
    add_normal(doc, "No disabled constraints data found.")

    add_normal(doc, _t(data, "Additional Suggestion:"), bold=True)
    add_normal(doc, "- You should verify invalid objects and disabled constraints that are used by application or not.")

    doc.add_page_break()


def add_appendix_defaults(doc, data, output_dir):
    add_heading1(doc, _t(data, "APPENDIX B – Information from alert log"))
    add_normal(doc, "- There is no alert Log to concern in this quarter.")
    doc.add_page_break()

    add_heading1(doc, _t(data, "APPENDIX C – SQL Statement should be to investigate"))
    add_normal(doc, "- There is no SQL statement to concern in this quarter.")
    doc.add_page_break()

    add_heading1(doc, _t(data, "APPENDIX D – Operating system log"))
    add_normal(doc, "- There is no Operating System log to concern in this quarter.")
    doc.add_page_break()

    add_heading1(doc, _t(data, "APPENDIX E – Backup Configuration"))
    backup_info = data.get("backup_info", {})
    if backup_info.get("found"):
        title = backup_info.get("title", "")
        if title:
            add_normal(doc, title, bold=True)
        log_content = backup_info.get("log_content", "")
        add_monospaced_box(doc, log_content, font_size=7)
    else:
        add_normal(doc, "No backup configuration data found.")
    doc.add_page_break()

    add_heading1(doc, _t(data, "APPENDIX F – OS Performance Summary"))
    
    os_info = data.get("os_info", {})
    os_perf = data.get("os_perf", {})
    
    if os_perf.get("found"):
        hostname = os_info.get("hostname", "N/A")
        os_release = os_info.get("os_version") or os_info.get("os_release") or "N/A"
        cpu_count = os_info.get("cpu_count") or os_info.get("cpu_cores") or "N/A"
        cores = os_info.get("cores_per_socket", "")
        ram_gb = os_info.get("total_ram_gb") or os_info.get("physical_memory_gb") or "N/A"

        cpu_label = f"{cpu_count} CPUs"
        if cores:
            cpu_label += f" ({cores} cores)"

        summary_table = add_small_table(
            doc,
            ["Name", "Value"],
            [
                {"Name": "Hostname", "Value": hostname},
                {"Name": "OS Version", "Value": os_release},
                {"Name": "CPU COUNT", "Value": cpu_label},
            ],
        )
        summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        cpu_path, mem_path, disk_path = create_os_perf_charts(os_perf, output_dir)
        chart_paths = os_perf.get("chart_paths", {})
        
        if chart_paths.get("cpu_all") or cpu_path:
            add_heading2(doc, _t(data, "F.1 CPU Performance"))
            
            runq_samples = os_perf.get("runq_samples", [])
            runq_trend = os_perf.get("runq_trend", [])
            if runq_samples or runq_trend:
                add_heading3(doc, "Process Run Queue")
                if chart_paths.get("runq"):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(chart_paths["runq"], width=Inches(6.8))

                runq_values = [float(r.get("runq", 0) or 0) for r in runq_samples]
                max_runq = max(runq_values, default=max((r.get("max", 0) for r in runq_trend), default=0))
                avg_runq = sum(runq_values) / max(len(runq_values), 1) if runq_values else sum(r.get("avg", 0) for r in runq_trend) / max(len(runq_trend), 1)

                add_normal(doc, "The following snaps recorded very high run queue values:")
                runq_table = add_small_table(doc, ["Rate", "Percent(%)"], [
                    {"Rate": "High (>3)", "Percent(%)": f"{_rate_percent(runq_samples, 'runq', 3):.2f}"},
                    {"Rate": "Very High (>6)", "Percent(%)": f"{_rate_percent(runq_samples, 'runq', 6):.2f}"},
                ], header_fill=HEADER_BLUE)
                runq_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                try:
                    cpu_c = int(float(cpu_count))
                except ValueError:
                    cpu_c = 2

                if max_runq > cpu_c * 2:
                    add_normal(doc, f"The run queue value is very high, indicating a severe CPU bottleneck. With only {cpu_c} CPUs, the run queue should ideally remain below {cpu_c}. However, the graph shows that the queue frequently spikes and reaches a maximum of {max_runq:.2f}. This means there are often more processes waiting for a CPU than CPUs available, which leads to significant performance delays.")
                    add_normal(doc, f"The average value of {avg_runq:.2f} does not reflect the severity of the frequent peak loads. CPU resources are insufficient for the current workload and upgrading or optimizing CPU usage is strongly recommended.")
                else:
                    add_normal(doc, f"The run queue value is normal. Maximum run queue is {max_runq:.2f} with an average of {avg_runq:.2f}. CPU resources are sufficient for the current workload.")
            
            add_heading3(doc, "CPU Utilization")
            for key, heading in [
                ("cpu_all", None),
                ("cpu_sys", "CPU Utilization by System"),
                ("cpu_usr", "CPU Utilization by User"),
            ]:
                if heading:
                    add_heading3(doc, heading)
                if chart_paths.get(key):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(chart_paths[key], width=Inches(6.8))

            cpu_samples = os_perf.get("cpu_samples", [])
            if cpu_samples:
                add_normal(doc, "The following recorded high CPU utilization values:")
                cpu_rate_table = add_small_table(doc, ["Rate", "Percent(%)"], [
                    {"Rate": "High (>95%)", "Percent(%)": f"{_rate_percent(cpu_samples, 'util', 95):.2f}"},
                    {"Rate": "Very High (100%)", "Percent(%)": f"{_rate_percent(cpu_samples, 'util', 99.99):.2f}"},
                ], header_fill=HEADER_BLUE)
                cpu_rate_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            cpu_trend = os_perf.get("cpu_trend", [])
            usr_vals = [r["usr"] for r in cpu_trend] or [r["usr"] for r in cpu_samples]
            sys_vals = [r["sys"] for r in cpu_trend] or [r["sys"] for r in cpu_samples]
            iow_vals = [r["iowait"] for r in cpu_trend] or [r["iowait"] for r in cpu_samples]
            
            sample_util = [r.get("util", 0) for r in cpu_samples]
            max_cpu = max(sample_util, default=max([u + s for u, s in zip(usr_vals, sys_vals)], default=0))
            avg_cpu = sum(sample_util) / max(len(sample_util), 1) if sample_util else sum([u + s for u, s in zip(usr_vals, sys_vals)]) / max(len(usr_vals), 1)
            
            if max_cpu >= 95:
                add_normal(doc, f"CPU utilization is high. The graph shows that CPU usage frequently reaches and sustains near 100% for extended periods. Highest CPU usage is {max_cpu:.2f}% while the average CPU usage is around {avg_cpu:.2f}%. The higher the CPU utilization, the longer it will take processes to run, indicating that the system is operating at full capacity and may experience performance bottlenecks.")
            else:
                add_normal(doc, f"CPU utilization is within acceptable limits. Highest CPU usage is {max_cpu:.2f}% while the average CPU usage is around {avg_cpu:.2f}%.")
            
            add_small_table(doc, ["Metric", "Min (%)", "Max (%)", "Avg (%)"], [
                {"Metric": "CPU Utilization by User (%usr)", "Min (%)": f"{min(usr_vals):.2f}", "Max (%)": f"{max(usr_vals):.2f}", "Avg (%)": f"{(sum(usr_vals)/len(usr_vals)):.2f}"},
                {"Metric": "CPU Utilization by System (%sys)", "Min (%)": f"{min(sys_vals):.2f}", "Max (%)": f"{max(sys_vals):.2f}", "Avg (%)": f"{(sum(sys_vals)/len(sys_vals)):.2f}"},
                {"Metric": "CPU Wait I/O (%iowait)", "Min (%)": f"{min(iow_vals):.2f}", "Max (%)": f"{max(iow_vals):.2f}", "Avg (%)": f"{(sum(iow_vals)/len(iow_vals)):.2f}"},
            ])
            
        if chart_paths.get("mem_free") or mem_path:
            add_heading2(doc, _t(data, "F.2 Memory Performance"))
            add_heading3(doc, "Physical Memory")
            if chart_paths.get("mem_free"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["mem_free"], width=Inches(6.8))

            add_heading3(doc, "Swap")
            if chart_paths.get("swap_used"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["swap_used"], width=Inches(6.8))

            add_heading3(doc, "Page in")
            if chart_paths.get("page_in"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["page_in"], width=Inches(6.8))

            add_heading3(doc, "Page out")
            if chart_paths.get("page_out"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["page_out"], width=Inches(6.8))

            total_ram = os_info.get("total_ram_mb", "")
            swap_total = os_info.get("swap_total_mb", "")
            if total_ram or swap_total:
                add_normal(doc, "System configuration:", bold=True)
                if total_ram:
                    try:
                        add_normal(doc, f"Total memory = {_fmt_number(float(total_ram), 2)} Megabytes ({_fmt_number(float(total_ram) / 1024, 2)} Gigabytes)")
                    except ValueError:
                        add_normal(doc, f"Total memory = {total_ram} Megabytes")
                if swap_total:
                    try:
                        add_normal(doc, f"Total swap = {_fmt_number(float(swap_total), 2)} Megabytes ({_fmt_number(float(swap_total) / 1024, 2)} Gigabytes)")
                    except ValueError:
                        add_normal(doc, f"Total swap = {swap_total} Megabytes")
            
            free_trend = os_perf.get("free_mem_trend", [])
            swap_trend = os_perf.get("swap_trend", [])
            mem_free_samples = os_perf.get("mem_free_samples", [])
            swap_samples = os_perf.get("swap_samples", [])
            
            if free_trend and swap_trend:
                min_free = min((r.get("free_gb", 0) for r in mem_free_samples), default=0)
                avg_free = sum(r.get("free_gb", 0) for r in mem_free_samples) / max(len(mem_free_samples), 1)
                max_swap_gb = max((r.get("used_gb", 0) for r in swap_samples), default=0)
                max_swap_pct = max((r.get("used_pct", 0) for r in swap_samples), default=max((r.get("max_pct", 0) for r in swap_trend), default=0))
                
                if max_swap_pct > 50.0 or min_free < 0.5:
                    add_normal(doc, f"At the highest workload, the free memory remains around {min_free:.2f} GB. This indicates that the current workload is close to or in excess of available physical memory, causing the system to use swap and resulting in degraded performance.")
                    add_normal(doc, f"The swap memory load reached around {max_swap_gb:.2f} GB. In some moments of time, high memory page in and page out rates are observed. That means OS is low on free memory. We strongly recommend consulting with the system team to monitor memory usage trends, investigate memory-intensive processes, and plan for a memory capacity upgrade.")
                    add_normal(doc, f"As a result of the above graph, memory should be monitored closely. Average free memory is around {avg_free:.2f} GB, and swap usage is significant under peak workload.")
                else:
                    add_normal(doc, f"Memory is sufficient for the current workload. At high workload, free memory remains around {min_free:.2f} GB and swap usage peak is {max_swap_pct:.2f}%. No immediate memory capacity upgrade is required.")
            
            mem_trend = os_perf.get("mem_trend", [])
            pgpgin_vals = [r["pgpgin"] for r in mem_trend]
            pgpgout_vals = [r["pgpgout"] for r in mem_trend]
            
            add_small_table(doc, ["Metric", "Min (Pages/s)", "Max (Pages/s)", "Avg (Pages/s)"], [
                {"Metric": "Page In (pgpgin/s)", "Min (Pages/s)": f"{min(pgpgin_vals):.2f}", "Max (Pages/s)": f"{max(pgpgin_vals):.2f}", "Avg (Pages/s)": f"{(sum(pgpgin_vals)/len(pgpgin_vals)):.2f}"},
                {"Metric": "Page Out (pgpgout/s)", "Min (Pages/s)": f"{min(pgpgout_vals):.2f}", "Max (Pages/s)": f"{max(pgpgout_vals):.2f}", "Avg (Pages/s)": f"{(sum(pgpgout_vals)/len(pgpgout_vals)):.2f}"},
            ])
            
        if chart_paths.get("cpu_wait_io") or chart_paths.get("disk_service") or disk_path:
            add_heading2(doc, _t(data, "F.3 Disk Performance"))
            add_heading3(doc, "CPU Wait I/O")
            if chart_paths.get("cpu_wait_io"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["cpu_wait_io"], width=Inches(6.8))

            add_heading3(doc, "Disk Service Time")
            if chart_paths.get("disk_service"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_paths["disk_service"], width=Inches(6.8))
            
            disk_trend = os_perf.get("disk_trend", [])
            disk_samples = _max_disk_samples_by_time(os_perf.get("disk_samples", []))
            svctm_vals = [r["svctm"] for r in disk_samples] or [r["svctm"] for r in disk_trend]
            await_vals = [r["await"] for r in disk_samples] or [r["await"] for r in disk_trend]
            iowait_vals = [r["iowait"] for r in os_perf.get("cpu_samples", [])]
            
            max_await = max(await_vals, default=0)
            max_iowait = max(iowait_vals, default=0)
            if max_iowait > 30 or max_await > 100:
                add_normal(doc, f"As the result from above graph, CPU Wait I/O occasionally peaks at high values. This suggests that system performance may still be impacted by high I/O demand or insufficient memory during peak workload.")
            else:
                add_normal(doc, "As the result from the graph, the disk service time and I/O wait time are generally low, indicating that disk devices are responding quickly. There are no continuous periods of high disk service time.")
            add_normal(doc, "In summary, disk performance is sufficient, but overall system performance may be degraded during periods of heavy memory pressure. Monitoring memory expansion may help reduce high I/O wait times.")
            
            add_small_table(doc, ["Metric", "Min (ms)", "Max (ms)", "Avg (ms)"], [
                {"Metric": "Service Time (svctm)", "Min (ms)": f"{min(svctm_vals):.2f}", "Max (ms)": f"{max(svctm_vals):.2f}", "Avg (ms)": f"{(sum(svctm_vals)/len(svctm_vals)):.2f}"},
                {"Metric": "I/O Wait (await)", "Min (ms)": f"{min(await_vals):.2f}", "Max (ms)": f"{max(await_vals):.2f}", "Avg (ms)": f"{(sum(await_vals)/len(await_vals)):.2f}"},
            ])
            
    else:
        add_normal(doc, "No OS performance summary data found.")


# =========================================================
# Main entry
# =========================================================

def create_docx_report(data: dict, output_docx: Path, source_zip: Path, output_dir: Path = None):
    output_docx = Path(output_docx)
    source_zip = Path(source_zip)
    output_dir = Path(output_dir) if output_dir else output_docx.parent

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    set_doc_layout(doc)
    set_update_fields_on_open(doc)
    add_report_header(doc, data)
    add_report_footer(doc, data)

    add_cover_page(doc, data, source_zip, output_dir)
    add_table_of_contents(doc, data)
    add_defect_classification(doc, data)
    add_suggestion_summary(doc, data)
    add_general_project_information(doc, data, source_zip)
    add_oracle_minimum_requirement(doc, data)
    add_system_checklist(doc, data)
    add_database_information(doc, data)
    add_database_growth_rate(doc, data, output_dir)
    add_performance_analysis(doc, data, output_dir)
    add_tablespace_free_space(doc, data)
    add_default_tablespace(doc, data)
    add_database_registry(doc, data)
    add_appendix_invalid_objects(doc, data)
    add_appendix_defaults(doc, data, output_dir)

    doc.save(output_docx)
